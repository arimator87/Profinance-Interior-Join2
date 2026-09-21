#!/usr/bin/env python3
"""
Additional verification: Check ZIP contents in detail
"""
import sys
import json
import zipfile
import io
import requests
from pathlib import Path
from openpyxl import load_workbook

# Read backend URL
env_path = Path("/app/frontend/.env")
BACKEND_URL = None
if env_path.exists():
    with open(env_path) as f:
        for line in f:
            if line.startswith("REACT_APP_BACKEND_URL="):
                BACKEND_URL = line.split("=", 1)[1].strip()
                break

BASE_URL = f"{BACKEND_URL}/api"

# Get demo token
response = requests.post(f"{BASE_URL}/auth/demo", timeout=10)
token = response.json()["token"]

# Get backup export
headers = {"Authorization": f"Bearer {token}"}
response = requests.get(f"{BASE_URL}/backup/export", headers=headers, timeout=30)

# Extract and analyze ZIP
zip_file = zipfile.ZipFile(io.BytesIO(response.content))

print("=" * 70)
print("📦 Detailed ZIP Content Analysis")
print("=" * 70)

# 1. Analyze data.json
print("\n1️⃣  data.json Analysis:")
data_json = json.loads(zip_file.read("data.json"))
print(f"   ✅ Valid JSON")
print(f"   📊 Keys: {list(data_json.keys())}")
print(f"   📁 Projects: {len(data_json.get('projects', []))}")
print(f"   💰 Transactions: {len(data_json.get('transactions', []))}")
print(f"   👷 Workers: {len(data_json.get('workers', []))}")
print(f"   📋 Work Items: {len(data_json.get('work_items', []))}")
print(f"   📈 Progress Entries: {len(data_json.get('progress_entries', []))}")
print(f"   📎 Files: {len(data_json.get('files', []))}")
print(f"   🛒 Orders: {len(data_json.get('orders', []))}")

# Check user info
user_info = data_json.get('user', {})
print(f"\n   👤 User Info:")
print(f"      - Email: {user_info.get('email')}")
print(f"      - Name: {user_info.get('name')}")
print(f"      - User ID: {user_info.get('user_id')}")

# Check exportedAt
print(f"\n   📅 Exported At: {data_json.get('exportedAt')}")

# Check photoMapping
photo_mapping = data_json.get('photoMapping', {})
print(f"\n   🖼️  Photo Mapping: {len(photo_mapping)} entries")
if photo_mapping:
    print(f"      Sample: {list(photo_mapping.items())[:2]}")

# Check sample project
if data_json.get('projects'):
    proj = data_json['projects'][0]
    print(f"\n   📁 Sample Project:")
    print(f"      - Name: {proj.get('name')}")
    print(f"      - Owner: {proj.get('owner')}")
    print(f"      - Nominal: Rp {proj.get('nominal'):,}")
    print(f"      - Status: {proj.get('status')}")

# 2. Analyze data.xlsx
print("\n2️⃣  data.xlsx Analysis:")
xlsx_bytes = zip_file.read("data.xlsx")
wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
print(f"   ✅ Valid Excel file")
print(f"   📊 Sheets: {wb.sheetnames}")

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    print(f"\n   📄 Sheet '{sheet_name}':")
    print(f"      - Rows: {len(rows)}")
    if rows:
        print(f"      - Headers: {rows[0]}")
        if len(rows) > 1:
            print(f"      - Sample data: {rows[1]}")

# 3. Analyze manifest.json
print("\n3️⃣  manifest.json Analysis:")
manifest = json.loads(zip_file.read("manifest.json"))
print(f"   ✅ Valid JSON")
print(f"   📱 App: {manifest.get('app')}")
print(f"   📅 Exported At: {manifest.get('exportedAt')}")
print(f"   📊 Counts:")
for key, value in manifest.get('counts', {}).items():
    print(f"      - {key}: {value}")
print(f"   📝 Notes: {manifest.get('notes')}")

# 4. Check all entries
print("\n4️⃣  All ZIP Entries:")
for entry in zip_file.namelist():
    info = zip_file.getinfo(entry)
    print(f"   - {entry} ({info.file_size} bytes)")

print("\n" + "=" * 70)
print("✅ Detailed analysis complete")
print("=" * 70)

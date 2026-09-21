"""Backup / export / restore module for ProFinance Interior.

Build a full ZIP archive of a user's data (or a single project):
  - data.json      : raw dump of all collections (restore-friendly) + photoMapping
  - data.xlsx      : human-readable spreadsheet (per-sheet)
  - photos/        : every referenced receipt, progress & thumbnail photo
  - manifest.json  : summary + counts

Photos stored on Emergent Object Storage are fetched via storage.get_object.
External http(s) photos (e.g. demo seed) are referenced by URL only (skipped).

Also supports:
  - create_stored_backup(): assemble ZIP and persist it to Object Storage + `backups` collection (auto/manual), with retention.
  - restore_from_zip(): import a backup ZIP, re-upload photos and restore missing projects + their records.
"""
import os
import io
import json
import uuid
import zipfile
import tempfile
import logging
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from auth import db
from storage import get_object, put_object, delete_object, APP_NAME, MIME_TYPES

logger = logging.getLogger(__name__)

RETENTION = 8  # keep this many stored backups per user


def _now_stamp():
    return datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")


def now_iso():
    return datetime.now(timezone.utc).isoformat()


async def gather_user_data(user: dict, project_ids=None) -> dict:
    """Collect documents belonging to the user. If project_ids given, scope to those projects."""
    uid = user["user_id"]
    proj_q = {"user_id": uid}
    if project_ids is not None:
        proj_q = {"user_id": uid, "id": {"$in": project_ids}}
    projects = await db.projects.find(proj_q, {"_id": 0}).to_list(2000)
    pids = [p["id"] for p in projects]
    transactions = await db.transactions.find({"project_id": {"$in": pids}}, {"_id": 0}).to_list(50000)
    workers = await db.workers.find({"project_id": {"$in": pids}}, {"_id": 0}).to_list(10000)
    work_items = await db.work_items.find({"project_id": {"$in": pids}}, {"_id": 0}).to_list(20000)
    sub_items = await db.sub_items.find({"project_id": {"$in": pids}}, {"_id": 0}).to_list(50000)
    wiids = [w["id"] for w in work_items]
    progress = await db.progress_entries.find({"workItemId": {"$in": wiids}}, {"_id": 0}).to_list(50000)
    # orders only meaningful for a full-account backup
    orders = []
    if project_ids is None:
        orders = await db.orders.find(
            {"user_id": uid}, {"_id": 0, "snap_token": 0, "redirect_url": 0}
        ).to_list(2000)
    # file metadata (for original filenames); all user files is fine as a lookup map
    files = await db.files.find({"user_id": uid}, {"_id": 0}).to_list(50000)
    return {
        "projects": projects,
        "transactions": transactions,
        "workers": workers,
        "work_items": work_items,
        "sub_items": sub_items,
        "progress_entries": progress,
        "files": files,
        "orders": orders,
    }


def _collect_photo_paths(data: dict) -> list:
    """Return unique object-storage paths referenced by the data (skip external http URLs)."""
    paths = set()

    def add(u):
        if u and isinstance(u, str) and not u.startswith("http"):
            paths.add(u.lstrip("/"))

    for p in data["projects"]:
        add(p.get("thumbnail"))
    for tx in data["transactions"]:
        add(tx.get("receiptUrl"))
    for e in data["progress_entries"]:
        for u in (e.get("photoUrls") or []):
            add(u)
    return sorted(paths)


def _build_xlsx_bytes(data: dict, user: dict) -> bytes:
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="EA580C")

    def style_header(ws):
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

    proj_name = {p["id"]: p.get("name", "") for p in data["projects"]}
    wi_name = {w["id"]: w.get("name", "") for w in data["work_items"]}
    wi_project = {w["id"]: w.get("project_id") for w in data["work_items"]}

    ws = wb.active
    ws.title = "Proyek"
    ws.append(["Nama Proyek", "Klien", "Perusahaan", "Nilai Kontrak (Rp)", "Kategori",
               "Status", "Alamat", "Tanggal Mulai", "Target Selesai", "Dibuat"])
    for p in data["projects"]:
        ws.append([
            p.get("name", ""), p.get("owner", ""), p.get("companyName", ""),
            p.get("nominal", 0), p.get("category", ""), p.get("status", ""),
            p.get("alamatProyek", ""), p.get("tanggalMulai", ""),
            p.get("targetSelesai", ""), p.get("createdAt", ""),
        ])
    style_header(ws)

    ws = wb.create_sheet("Transaksi")
    ws.append(["Proyek", "Tipe", "Kategori", "Deskripsi", "Nominal (Rp)", "Tanggal", "Struk"])
    for tx in data["transactions"]:
        rec = tx.get("receiptUrl")
        struk = "Foto tersimpan" if rec and not str(rec).startswith("http") else ("URL" if rec else "-")
        ws.append([
            proj_name.get(tx.get("project_id"), ""),
            "Masuk" if tx.get("type") == "in" else "Keluar",
            tx.get("category", ""), tx.get("description", ""),
            tx.get("amount", 0), tx.get("date", ""), struk,
        ])
    style_header(ws)

    ws = wb.create_sheet("Tukang")
    ws.append(["Proyek", "Nama Tukang", "Borongan (Rp)"])
    for w in data["workers"]:
        ws.append([proj_name.get(w.get("project_id"), ""), w.get("name", ""), w.get("borongan", 0)])
    style_header(ws)

    ws = wb.create_sheet("Item Pekerjaan")
    ws.append(["Proyek", "Item Pekerjaan", "Nilai (Rp)", "Sub Item", "Harga Sub (Rp)", "Selesai"])
    subs_by_wi = {}
    for s in data["sub_items"]:
        subs_by_wi.setdefault(s.get("work_item_id"), []).append(s)
    for w in data["work_items"]:
        subs = subs_by_wi.get(w["id"], [])
        if subs:
            for s in subs:
                ws.append([proj_name.get(w.get("project_id"), ""), w.get("name", ""),
                           w.get("nilai", 0), s.get("name", ""), s.get("harga", 0),
                           "Ya" if s.get("status") else "Tidak"])
        else:
            ws.append([proj_name.get(w.get("project_id"), ""), w.get("name", ""),
                       w.get("nilai", 0), "-", "", ""])
    style_header(ws)

    ws = wb.create_sheet("Progress Lapangan")
    ws.append(["Proyek", "Item Pekerjaan", "Tanggal", "Progress (%)", "Catatan", "Jumlah Foto"])
    for e in data["progress_entries"]:
        wi = e.get("workItemId")
        ws.append([
            proj_name.get(wi_project.get(wi), ""), wi_name.get(wi, ""),
            e.get("date", ""), e.get("progress", 0), e.get("notes", ""),
            len(e.get("photoUrls") or []),
        ])
    style_header(ws)

    for sheet in wb.worksheets:
        for col in sheet.columns:
            length = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(length + 2, 10), 45)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(tmp.name)
    tmp.close()
    with open(tmp.name, "rb") as fh:
        content = fh.read()
    os.unlink(tmp.name)
    return content


def _assemble_zip(data: dict, user: dict, photo_paths: list, scope_label: str = "Semua Proyek") -> tuple:
    """Synchronous: build xlsx, fetch photos, write zip to a temp file. Returns (path, stats)."""
    files_map = {f.get("storage_path"): f for f in data["files"]}
    stamp = _now_stamp()
    tmp_zip = tempfile.NamedTemporaryFile(delete=False, suffix=".zip")
    tmp_zip.close()

    photo_ok, photo_fail = 0, 0
    manifest_photos = {}
    used_names = set()

    with zipfile.ZipFile(tmp_zip.name, "w", zipfile.ZIP_DEFLATED) as zf:
        try:
            zf.writestr("data.xlsx", _build_xlsx_bytes(data, user))
        except Exception as e:
            logger.error(f"xlsx build failed: {e}")

        for path in photo_paths:
            meta = files_map.get(path, {})
            base = meta.get("original_filename") or path.split("/")[-1]
            base = base.replace("/", "_").replace("\\", "_")
            name = base
            i = 1
            while name in used_names:
                stem, dot, ext = base.rpartition(".")
                name = f"{stem}_{i}.{ext}" if dot else f"{base}_{i}"
                i += 1
            used_names.add(name)
            try:
                content, _ct = get_object(path)
                zf.writestr(f"photos/{name}", content)
                manifest_photos[path] = f"photos/{name}"
                photo_ok += 1
            except Exception as e:
                logger.warning(f"backup photo fetch failed {path}: {e}")
                manifest_photos[path] = None
                photo_fail += 1

        dump = {
            "exportedAt": now_iso(),
            "scope": scope_label,
            "user": {
                "user_id": user.get("user_id"),
                "email": user.get("email"),
                "name": user.get("name"),
            },
            **data,
            "photoMapping": manifest_photos,
        }
        zf.writestr("data.json", json.dumps(dump, ensure_ascii=False, indent=2, default=str))

        manifest = {
            "app": "ProFinance Interior",
            "exportedAt": dump["exportedAt"],
            "scope": scope_label,
            "counts": {
                "projects": len(data["projects"]),
                "transactions": len(data["transactions"]),
                "workers": len(data["workers"]),
                "work_items": len(data["work_items"]),
                "progress_entries": len(data["progress_entries"]),
                "photos_included": photo_ok,
                "photos_failed": photo_fail,
            },
            "notes": "Foto di folder photos/. data.json berisi data mentah untuk pemulihan.",
        }
        zf.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))

    stats = {"photos_included": photo_ok, "photos_failed": photo_fail,
             "projects": len(data["projects"]),
             "transactions": len(data["transactions"]), "stamp": stamp}
    return tmp_zip.name, stats


# ---------- Stored backups (auto/manual) ----------
async def create_stored_backup(user: dict, kind: str = "auto") -> dict:
    """Assemble a full backup and persist it to Object Storage + `backups` collection."""
    from fastapi.concurrency import run_in_threadpool
    uid = user["user_id"]
    data = await gather_user_data(user)
    if not data["projects"]:
        return None  # nothing to back up
    photo_paths = _collect_photo_paths(data)
    zip_path, stats = await run_in_threadpool(_assemble_zip, data, user, photo_paths, "Semua Proyek")
    with open(zip_path, "rb") as fh:
        content = fh.read()
    try:
        os.unlink(zip_path)
    except Exception:
        pass

    stamp = stats["stamp"]
    storage_path = f"{APP_NAME}/backups/{uid}/{stamp}.zip"
    await run_in_threadpool(put_object, storage_path, content, "application/zip")

    doc = {
        "id": str(uuid.uuid4()),
        "user_id": uid,
        "storage_path": storage_path,
        "filename": f"profinance-backup-{stamp}.zip",
        "size": len(content),
        "kind": kind,
        "counts": {
            "projects": stats["projects"],
            "transactions": stats["transactions"],
            "photos": stats["photos_included"],
        },
        "createdAt": now_iso(),
    }
    await db.backups.insert_one(doc)
    doc.pop("_id", None)

    # retention: keep newest RETENTION, delete older objects + docs
    old = await db.backups.find({"user_id": uid}, {"_id": 0}).sort("createdAt", -1).to_list(1000)
    for stale in old[RETENTION:]:
        await run_in_threadpool(delete_object, stale["storage_path"])
        await db.backups.delete_one({"id": stale["id"]})
    return doc


async def run_weekly_backups() -> dict:
    """Cron job: create a stored backup for every eligible (non-demo) user with projects."""
    created, skipped = 0, 0
    users = await db.users.find(
        {"isDemo": {"$ne": True}}, {"_id": 0, "user_id": 1, "email": 1, "name": 1}
    ).to_list(100000)
    for u in users:
        try:
            res = await create_stored_backup(u, kind="auto")
            if res:
                created += 1
            else:
                skipped += 1
        except Exception as e:
            logger.error(f"weekly backup failed for {u.get('email')}: {e}")
            skipped += 1
    logger.info(f"weekly backup done: created={created} skipped={skipped}")
    return {"created": created, "skipped": skipped}


# ---------- Restore ----------
async def restore_from_zip(user: dict, zip_bytes: bytes) -> dict:
    """Import a backup ZIP. Re-upload photos and restore projects that don't already exist."""
    from fastapi.concurrency import run_in_threadpool
    uid = user["user_id"]

    try:
        zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
    except Exception:
        raise ValueError("File bukan ZIP yang valid.")
    names = set(zf.namelist())
    if "data.json" not in names:
        raise ValueError("ZIP tidak berisi data.json (bukan backup ProFinance yang valid).")
    try:
        data = json.loads(zf.read("data.json").decode("utf-8"))
    except Exception:
        raise ValueError("data.json rusak / tidak dapat dibaca.")

    photo_mapping = data.get("photoMapping", {}) or {}

    # Re-upload photos referenced in the archive -> remap old_path -> new canonical path
    remap = {}
    photos_restored = 0
    for old_path, zip_name in photo_mapping.items():
        if not zip_name or zip_name not in names:
            remap[old_path] = old_path  # keep original reference (object may still exist)
            continue
        try:
            content = zf.read(zip_name)
        except Exception:
            remap[old_path] = old_path
            continue
        ext = old_path.rsplit(".", 1)[-1].lower() if "." in old_path else "bin"
        ct = MIME_TYPES.get(ext, "application/octet-stream")
        new_path = f"{APP_NAME}/uploads/{uid}/{uuid.uuid4()}.{ext}"
        try:
            result = await run_in_threadpool(put_object, new_path, content, ct)
            canonical = result.get("path", new_path)
            remap[old_path] = canonical
            await db.files.insert_one({
                "id": str(uuid.uuid4()),
                "user_id": uid,
                "storage_path": canonical,
                "original_filename": zip_name.split("/")[-1],
                "content_type": ct,
                "size": len(content),
                "is_deleted": False,
                "created_at": now_iso(),
            })
            photos_restored += 1
        except Exception as e:
            logger.warning(f"restore photo upload failed {old_path}: {e}")
            remap[old_path] = old_path

    def rp(u):
        if u and isinstance(u, str) and not u.startswith("http"):
            return remap.get(u.lstrip("/"), u)
        return u

    # Restore projects that don't already exist for this user
    existing = await db.projects.find({"user_id": uid}, {"_id": 0, "id": 1}).to_list(5000)
    existing_ids = {p["id"] for p in existing}

    restored_pids = set()
    restored_projects = 0
    skipped_projects = 0
    for p in data.get("projects", []):
        pid = p.get("id")
        if not pid or pid in existing_ids:
            skipped_projects += 1
            continue
        doc = {k: v for k, v in p.items() if k != "_id"}
        doc["user_id"] = uid
        doc["thumbnail"] = rp(doc.get("thumbnail"))
        # fresh portal token suffix to avoid slug clashes
        base_slug = (doc.get("portalSlug") or pid).rsplit("-", 1)[0]
        doc["portalSlug"] = f"{base_slug}-{uuid.uuid4().hex[:6]}"
        await db.projects.insert_one(doc)
        restored_pids.add(pid)
        restored_projects += 1

    if not restored_pids:
        return {
            "restored_projects": 0,
            "skipped_projects": skipped_projects,
            "photos_restored": photos_restored,
            "restored_transactions": 0,
        }

    # Transactions
    tx_count = 0
    for tx in data.get("transactions", []):
        if tx.get("project_id") in restored_pids:
            doc = {k: v for k, v in tx.items() if k != "_id"}
            doc["receiptUrl"] = rp(doc.get("receiptUrl"))
            await db.transactions.insert_one(doc)
            tx_count += 1

    # Workers
    for w in data.get("workers", []):
        if w.get("project_id") in restored_pids:
            await db.workers.insert_one({k: v for k, v in w.items() if k != "_id"})

    # Work items
    restored_wiids = set()
    for wi in data.get("work_items", []):
        if wi.get("project_id") in restored_pids:
            doc = {k: v for k, v in wi.items() if k != "_id"}
            await db.work_items.insert_one(doc)
            restored_wiids.add(wi.get("id"))

    # Sub items
    for s in data.get("sub_items", []):
        if s.get("project_id") in restored_pids or s.get("work_item_id") in restored_wiids:
            await db.sub_items.insert_one({k: v for k, v in s.items() if k != "_id"})

    # Progress entries (remap photoUrls)
    for e in data.get("progress_entries", []):
        if e.get("workItemId") in restored_wiids:
            doc = {k: v for k, v in e.items() if k != "_id"}
            doc["photoUrls"] = [rp(u) for u in (doc.get("photoUrls") or [])]
            await db.progress_entries.insert_one(doc)

    return {
        "restored_projects": restored_projects,
        "skipped_projects": skipped_projects,
        "photos_restored": photos_restored,
        "restored_transactions": tx_count,
    }

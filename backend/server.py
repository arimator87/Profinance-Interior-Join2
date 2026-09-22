import os
import re
import io
import html
import json
import hmac
import base64
import hashlib
import uuid
import secrets
import logging
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Optional, List, Literal
from urllib.parse import quote

from openpyxl import load_workbook, Workbook
import httpx

from fastapi import (
    FastAPI, APIRouter, Depends, HTTPException, Request, Response,
    UploadFile, File, Header, Query, BackgroundTasks,
)
from fastapi.responses import StreamingResponse, FileResponse
from fastapi.concurrency import run_in_threadpool
from starlette.background import BackgroundTask
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from pydantic import BaseModel, EmailStr, Field

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

import auth as auth_mod
from auth import (
    db, get_current_user, require_premium, create_session, set_session_cookie,
    register_email_user, login_email_user, process_google_session, user_public, OWNER_EMAILS,
    reset_password_with_phone,
)
from storage import init_storage, put_object, get_object, APP_NAME, MIME_TYPES
from pdf_report import build_report_pdf, build_progress_pdf, build_rab_pdf, build_invoice_pdf
import backup as backup_mod

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(name)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

app = FastAPI()
api = APIRouter(prefix="/api")

PAID_IN_CATEGORIES = {"Downpayment", "Termin", "Pelunasan"}
INCOME_CATEGORIES = ["Downpayment", "Termin", "Pelunasan", "Lainnya"]
EXPENSE_CATEGORIES = ["Material", "Makan", "Toll", "Bensin", "Lainnya"]


def now_iso():
    return datetime.now(timezone.utc).isoformat()


DEMO_PHOTOS = [
    "https://images.unsplash.com/photo-1736182615481-3795ea557614?crop=entropy&cs=srgb&fm=jpg&q=85&w=1080",
    "https://images.unsplash.com/photo-1618832515490-e181c4794a45?crop=entropy&cs=srgb&fm=jpg&q=85&w=1080",
    "https://images.unsplash.com/photo-1692890659047-079b769ee3e6?crop=entropy&cs=srgb&fm=jpg&q=85&w=1080",
    "https://images.unsplash.com/photo-1543525324-26e03b510586?crop=entropy&cs=srgb&fm=jpg&q=85&w=1080",
    "https://images.pexels.com/photos/5691533/pexels-photo-5691533.jpeg?auto=compress&cs=tinysrgb&dpr=2&h=650&w=940",
    "https://images.pexels.com/photos/36035073/pexels-photo-36035073.jpeg?auto=compress&cs=tinysrgb&dpr=2&h=650&w=940",
    "https://images.pexels.com/photos/15124970/pexels-photo-15124970.jpeg?auto=compress&cs=tinysrgb&dpr=2&h=650&w=940",
    "https://images.unsplash.com/photo-1772442198620-3674b427e59e?crop=entropy&cs=srgb&fm=jpg&q=85&w=1080",
]


def slugify(name: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", (name or "proyek").lower()).strip("-")
    return s[:40] or "proyek"


async def ensure_portal_slug(p: dict) -> str:
    if p.get("portalSlug"):
        return p["portalSlug"]
    slug = f"{slugify(p.get('name'))}-{secrets.token_hex(3)}"
    await db.projects.update_one({"id": p["id"]}, {"$set": {"portalSlug": slug}})
    return slug


async def ensure_rab_slug(p: dict) -> str:
    if p.get("rabSlug"):
        return p["rabSlug"]
    slug = f"{slugify(p.get('name'))}-rab-{secrets.token_hex(3)}"
    await db.projects.update_one({"id": p["id"]}, {"$set": {"rabSlug": slug}})
    return slug


def compute_rab(rab: dict) -> dict:
    """Compute RAB totals. subItem.nilai = round(qty*hargaSatuan) + sum(material.nilai)."""
    rab = rab or {}
    sections_out = []
    total_items = 0
    idx = 0
    for sec in rab.get("sections", []) or []:
        subs_out = []
        subtotal = 0
        for si in sec.get("subItems", []) or []:
            idx += 1
            try:
                qty = float(si.get("qty") or 0)
            except (TypeError, ValueError):
                qty = 0.0
            harga = int(si.get("hargaSatuan") or 0)
            base = round(qty * harga)
            mats = si.get("materials", []) or []
            mat_total = sum(int(m.get("nilai") or 0) for m in mats)
            nilai = base + mat_total
            subtotal += nilai
            subs_out.append({
                "id": si.get("id"), "no": idx, "name": si.get("name", ""),
                "qty": qty, "unit": si.get("unit", "Unit"), "hargaSatuan": harga,
                "base": base, "materials": mats, "materialTotal": mat_total, "nilai": nilai,
            })
        total_items += subtotal
        sections_out.append({"id": sec.get("id"), "name": sec.get("name", ""), "subItems": subs_out, "subtotal": subtotal})
    discount = int(rab.get("discount") or 0)
    after = total_items - discount
    ppn_enabled = bool(rab.get("ppnEnabled"))
    ppn_percent = float(rab.get("ppnPercent") or 0)
    ppn_amount = round(after * ppn_percent / 100) if ppn_enabled else 0
    grand = after + ppn_amount
    termins = []
    for t in rab.get("termins", []) or []:
        pct = float(t.get("percent") or 0)
        termins.append({"label": t.get("label", ""), "percent": pct, "nominal": round(grand * pct / 100)})
    return {
        "sections": sections_out, "totalItems": total_items, "discount": discount,
        "afterDiscount": after, "ppnEnabled": ppn_enabled, "ppnPercent": ppn_percent,
        "ppnAmount": ppn_amount, "grandTotal": grand, "termins": termins,
    }


def compute_invoice(inv: dict) -> dict:
    """Financial computation for an invoice/proforma/retention document."""
    items_out = []
    subtotal = 0
    for i, it in enumerate(inv.get("items", []) or [], start=1):
        qty = float(it.get("qty") or 0)
        unit = int(it.get("unitPrice") or 0)
        amount = round(qty * unit)
        subtotal += amount
        items_out.append({
            "no": i, "description": it.get("description", ""),
            "qty": qty, "unitPrice": unit, "amount": amount,
        })
    ppn_enabled = bool(inv.get("ppnEnabled"))
    ppn_percent = float(inv.get("ppnPercent") or 0)
    ppn_amount = round(subtotal * ppn_percent / 100) if ppn_enabled else 0
    gross = subtotal + ppn_amount
    ret_enabled = bool(inv.get("retentionEnabled"))
    ret_percent = float(inv.get("retentionPercent") or 0)
    # Retention is held back on the work value (DPP / subtotal), per common practice
    ret_amount = round(subtotal * ret_percent / 100) if ret_enabled else 0
    amount_due = gross - ret_amount
    return {
        "items": items_out, "subtotal": subtotal,
        "ppnEnabled": ppn_enabled, "ppnPercent": ppn_percent, "ppnAmount": ppn_amount,
        "grossTotal": gross,
        "retentionEnabled": ret_enabled, "retentionPercent": ret_percent, "retentionAmount": ret_amount,
        "amountDue": amount_due,
    }



# ---------- Schemas ----------
class RegisterIn(BaseModel):
    email: EmailStr
    name: str
    password: str = Field(min_length=6)
    phone: Optional[str] = ""


class ResetPasswordIn(BaseModel):
    email: EmailStr
    phone: str
    newPassword: str = Field(min_length=6)


class LoginIn(BaseModel):
    email: EmailStr
    password: str


class GoogleSessionIn(BaseModel):
    session_id: str


class ProjectIn(BaseModel):
    name: str
    owner: str = ""
    nominal: int = 0
    companyName: str = ""
    alamatProyek: str = ""
    tanggalMulai: Optional[str] = None
    targetSelesai: Optional[str] = None
    category: str = "Residensial"
    status: str = "Berjalan"
    thumbnail: Optional[str] = None


class TransactionIn(BaseModel):
    type: str
    amount: int
    description: str = ""
    date: Optional[str] = None
    category: str = "Lainnya"
    receiptUrl: Optional[str] = None


class WorkerIn(BaseModel):
    name: str
    borongan: int = 0


class WorkerPayIn(BaseModel):
    type: str  # "kasbon" | "pelunasan"
    amount: int
    date: Optional[str] = None
    description: str = ""


class WorkItemIn(BaseModel):
    name: str
    nilai: int = 0
    startDate: Optional[str] = None
    endDate: Optional[str] = None


class ProgressIn(BaseModel):
    date: Optional[str] = None
    progress: int = 0
    notes: str = ""
    photoUrls: List[str] = []
    subItemId: Optional[str] = None


class SubItemIn(BaseModel):
    name: str
    harga: int = 0
    status: bool = False


class RabIn(BaseModel):
    rabTotal: int = 0


class RabMaterialIn(BaseModel):
    id: Optional[str] = None
    name: str = ""
    nilai: int = 0


class RabSubItemIn(BaseModel):
    id: Optional[str] = None
    name: str = ""
    qty: float = 1
    unit: str = "Unit"
    hargaSatuan: int = 0
    materials: List[RabMaterialIn] = []


class RabSectionIn(BaseModel):
    id: Optional[str] = None
    name: str = ""
    subItems: List[RabSubItemIn] = []


class RabTerminIn(BaseModel):
    label: str = ""
    percent: float = 0


class RabDoc(BaseModel):
    clientName: str = ""
    clientAddress: str = ""
    clientPhone: str = ""
    quotationNo: str = ""
    quotationDate: Optional[str] = None
    companyName: str = ""
    companyAddress: str = ""
    companyPhone: str = ""
    bankName: str = ""
    bankAccount: str = ""
    bankHolder: str = ""
    signerLeft: str = ""
    signerRight: str = ""
    signatureImage: str = ""
    notes: str = ""
    discount: int = 0
    ppnEnabled: bool = False
    ppnPercent: float = 11
    sections: List[RabSectionIn] = []
    termins: List[RabTerminIn] = []


class RabCreateIn(BaseModel):
    projectName: str
    category: str = "Residensial"
    rab: RabDoc


class CompanyTemplateIn(BaseModel):
    companyName: str = ""
    companyAddress: str = ""
    companyPhone: str = ""
    bankName: str = ""
    bankAccount: str = ""
    bankHolder: str = ""
    signerLeft: str = ""
    signerRight: str = ""
    signatureImage: str = ""


class InvoiceItemIn(BaseModel):
    description: str = ""
    qty: float = 1
    unitPrice: int = 0


class InvoiceIn(BaseModel):
    type: str = "proforma"  # proforma | final | retention
    number: str = ""
    quotationNo: str = ""
    invoiceDate: Optional[str] = None
    dueDate: Optional[str] = None
    status: str = "Draft"  # Draft | Terkirim | Lunas
    clientName: str = ""
    clientAddress: str = ""
    clientPhone: str = ""
    items: List[InvoiceItemIn] = []
    ppnEnabled: bool = False
    ppnPercent: float = 11
    retentionEnabled: bool = False
    retentionPercent: float = 5
    companyName: str = ""
    companyAddress: str = ""
    companyPhone: str = ""
    bankName: str = ""
    bankAccount: str = ""
    bankHolder: str = ""
    signerLeft: str = ""
    signerRight: str = ""
    signatureImage: str = ""
    notes: str = ""


# ---------- Finance calc ----------
async def compute_summary(project: dict) -> dict:
    project_id = project["id"]
    nominal = project.get("nominal", 0) or 0
    txs = await db.transactions.find({"project_id": project_id}, {"_id": 0}).to_list(5000)
    total_in = sum(t["amount"] for t in txs if t["type"] == "in")
    total_out = sum(t["amount"] for t in txs if t["type"] == "out")
    balance = total_in - total_out
    margin_pct = (balance / nominal * 100) if nominal else 0
    realisasi_pct = (total_in / nominal * 100) if nominal else 0
    terbayar = sum(t["amount"] for t in txs if t["type"] == "in" and t.get("category") in PAID_IN_CATEGORIES)
    sisa_tagihan = nominal - terbayar
    return {
        "totalIn": total_in,
        "totalOut": total_out,
        "balance": balance,
        "marginPct": round(margin_pct, 2),
        "realisasiPct": round(realisasi_pct, 2),
        "terbayar": terbayar,
        "sisaTagihan": sisa_tagihan,
        "nominal": nominal,
        "txCount": len(txs),
    }


async def compute_worker(worker: dict) -> dict:
    project_id = worker["project_id"]
    name = worker["name"]
    txs = await db.transactions.find(
        {"project_id": project_id, "type": "out"}, {"_id": 0}
    ).to_list(5000)
    total_kasbon = sum(t["amount"] for t in txs if t.get("category") == f"Kasbon Tukang {name}")
    total_pelunasan = sum(t["amount"] for t in txs if t.get("category") == f"Pelunasan Tukang {name}")
    sisa = worker.get("borongan", 0) - total_kasbon - total_pelunasan
    return {
        **{k: v for k, v in worker.items() if k != "project_id"},
        "totalKasbon": total_kasbon,
        "totalPelunasan": total_pelunasan,
        "totalDibayar": total_kasbon + total_pelunasan,
        "sisaHutang": sisa,
    }


async def compute_progress(project: dict):
    project_id = project["id"]
    items_raw = await db.work_items.find({"project_id": project_id}, {"_id": 0}).sort("createdAt", 1).to_list(2000)
    prepared = []
    total_item_value = 0
    for wi in items_raw:
        subs = await db.sub_items.find({"work_item_id": wi["id"]}, {"_id": 0}).sort("createdAt", 1).to_list(2000)
        has_subs = len(subs) > 0
        item_value = sum(s.get("harga", 0) for s in subs) if has_subs else (wi.get("nilai", 0) or 0)
        total_item_value += item_value
        prepared.append((wi, subs, has_subs, item_value))

    rab = project.get("rabTotal", 0) or 0
    if rab <= 0:
        rab = total_item_value

    async def last_prog(work_item_id, sub_item_id):
        q = {"workItemId": work_item_id}
        if sub_item_id is None:
            q["$or"] = [{"subItemId": None}, {"subItemId": {"$exists": False}}]
        else:
            q["subItemId"] = sub_item_id
        entries = await db.progress_entries.find(q, {"_id": 0}).sort("date", 1).to_list(2000)
        return (entries[-1]["progress"] if entries else 0), len(entries)

    items = []
    project_completed = 0
    for wi, subs, has_subs, item_value in prepared:
        item_weight = (item_value / rab * 100) if rab else 0
        sub_list = []
        if has_subs:
            completed = 0
            for s in subs:
                sp, scount = await last_prog(wi["id"], s["id"])
                sw = (s.get("harga", 0) / rab * 100) if rab else 0
                completed += s.get("harga", 0) * sp / 100
                sub_list.append({
                    "id": s["id"], "name": s["name"], "harga": s.get("harga", 0),
                    "weight": round(sw, 2), "lastProgress": sp, "entryCount": scount,
                })
            item_progress = round(completed / item_value * 100, 1) if item_value else 0
            item_completed = completed
            entry_count = sum(x["entryCount"] for x in sub_list)
        else:
            ip, icount = await last_prog(wi["id"], None)
            item_progress = ip
            item_completed = item_value * ip / 100
            entry_count = icount
        project_completed += item_completed
        items.append({
            "id": wi["id"], "name": wi["name"], "nilai": item_value, "manualNilai": wi.get("nilai", 0),
            "startDate": wi.get("startDate"), "endDate": wi.get("endDate"),
            "weight": round(item_weight, 2), "lastProgress": item_progress,
            "hasSubs": has_subs, "subItems": sub_list, "subCount": len(subs),
            "subTotal": item_value if has_subs else 0, "doneValue": round(item_completed),
            "entryCount": entry_count,
        })

    total_progress = round(project_completed / rab * 100, 2) if rab else 0
    return {
        "items": items, "rab": rab, "rabTotal": project.get("rabTotal", 0) or 0,
        "totalItemValue": total_item_value, "totalProgress": total_progress,
        "completedValue": round(project_completed),
        "rabBaseline": project.get("rabBaseline"),
    }


async def compute_workitems(project: dict):
    return (await compute_progress(project))["items"]


async def project_total_progress(project: dict):
    data = await compute_progress(project)
    return data["totalProgress"], data["items"]


async def get_owned_project(project_id: str, user: dict) -> dict:
    project = await db.projects.find_one({"id": project_id, "user_id": user["user_id"]}, {"_id": 0})
    if not project:
        raise HTTPException(status_code=404, detail="Proyek tidak ditemukan")
    return project


# ---------- Auth routes ----------
@api.post("/auth/register")
async def register(body: RegisterIn, response: Response):
    user = await register_email_user(body.email, body.name, body.password, body.phone or "")
    token = await create_session(user["user_id"])
    set_session_cookie(response, token)
    return {"user": user_public(user), "token": token}


@api.post("/auth/reset-password")
async def reset_password(body: ResetPasswordIn):
    await reset_password_with_phone(body.email, body.phone, body.newPassword)
    return {"ok": True, "message": "Password berhasil diperbarui. Silakan masuk dengan password baru."}


@api.post("/auth/login")
async def login(body: LoginIn, response: Response):
    user = await login_email_user(body.email, body.password)
    token = await create_session(user["user_id"])
    set_session_cookie(response, token)
    return {"user": user_public(user), "token": token}


@api.post("/auth/session")
async def google_session(body: GoogleSessionIn, response: Response):
    user, token = await process_google_session(body.session_id)
    set_session_cookie(response, token)
    return {"user": user_public(user), "token": token}


@api.post("/auth/demo")
async def demo_login(response: Response):
    demo_email = "demo@profinance.id"
    long_expiry = (datetime.now(timezone.utc) + timedelta(days=3650)).isoformat()
    user = await db.users.find_one({"email": demo_email}, {"_id": 0})
    if not user:
        user = {
            "user_id": f"user_demo_{uuid.uuid4().hex[:8]}",
            "email": demo_email, "name": "Akun Demo", "password_hash": None, "picture": None,
            "subscriptionTier": "premium", "subscriptionExpiry": long_expiry,
            "stripeCustomerId": None, "authProvider": "demo", "isDemo": True,
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.users.insert_one(user)
        user.pop("_id", None)
    await db.users.update_one(
        {"user_id": user["user_id"]},
        {"$set": {"subscriptionTier": "premium", "subscriptionExpiry": long_expiry, "isDemo": True}},
    )
    count = await db.projects.count_documents({"user_id": user["user_id"]})
    if count == 0:
        await _provision_demo_data(user["user_id"])
    token = await create_session(user["user_id"])
    set_session_cookie(response, token)
    user = await db.users.find_one({"user_id": user["user_id"]}, {"_id": 0})
    return {"user": user_public(user), "token": token}


async def _provision_demo_data(user_id: str):
    """Seed sample projects + a baseline showcase for the demo account (idempotent)."""
    if await db.projects.count_documents({"user_id": user_id}) > 0:
        return
    await seed_demo({"user_id": user_id})
    proj = await db.projects.find_one({"user_id": user_id}, {"_id": 0}, sort=[("createdAt", 1)])
    if not proj:
        return
    wis = await db.work_items.find({"project_id": proj["id"]}, {"_id": 0}).sort("createdAt", 1).to_list(50)
    if wis:
        baseline = {
            "savedAt": (datetime.now(timezone.utc) - timedelta(days=20)).isoformat(),
            "rabTotal": 0,
            "totalItemValue": sum(w.get("nilai", 0) for w in wis),
            "items": [{"id": w["id"], "name": w["name"], "value": w.get("nilai", 0)} for w in wis],
        }
        await db.work_items.update_one({"id": wis[0]["id"]}, {"$set": {"nilai": wis[0].get("nilai", 0) + 20000000}})
        await db.projects.update_one({"id": proj["id"]}, {"$set": {"rabBaseline": baseline}})


async def _reset_demo_job():
    user = await db.users.find_one({"email": "demo@profinance.id"}, {"_id": 0})
    if not user:
        return
    uid = user["user_id"]
    projects = await db.projects.find({"user_id": uid}, {"_id": 0, "id": 1, "name": 1, "portalSlug": 1}).to_list(1000)
    # Preserve existing portal slugs so client portal links already shared keep working after reset
    slug_by_name = {p["name"]: p["portalSlug"] for p in projects if p.get("portalSlug")}
    pids = [p["id"] for p in projects]
    witems = await db.work_items.find({"project_id": {"$in": pids}}, {"_id": 0, "id": 1}).to_list(3000)
    wiids = [w["id"] for w in witems]
    await db.progress_entries.delete_many({"workItemId": {"$in": wiids}})
    await db.sub_items.delete_many({"project_id": {"$in": pids}})
    await db.work_items.delete_many({"project_id": {"$in": pids}})
    await db.transactions.delete_many({"project_id": {"$in": pids}})
    await db.workers.delete_many({"project_id": {"$in": pids}})
    await db.projects.delete_many({"user_id": uid})
    await _provision_demo_data(uid)
    # Re-apply preserved slugs onto the freshly seeded projects (matched by name)
    new_projects = await db.projects.find({"user_id": uid}, {"_id": 0, "id": 1, "name": 1, "portalSlug": 1}).to_list(1000)
    for np in new_projects:
        slug = slug_by_name.get(np["name"]) or f"{slugify(np['name'])}-demo"
        if np.get("portalSlug") != slug:
            await db.projects.update_one({"id": np["id"]}, {"$set": {"portalSlug": slug}})
    logger.info("Demo data reset & re-seeded for %s (portal slugs preserved: %d)", uid, len(slug_by_name))


@api.post("/cron/reset-demo")
async def cron_reset_demo(request: Request, background: BackgroundTasks):
    # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
    secret = os.environ.get("WEBHOOK_CRON_SECRET")
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else ""
    if not secret or not hmac.compare_digest(token, secret):
        raise HTTPException(status_code=401, detail="Unauthorized")
    background.add_task(_reset_demo_job)
    return {"ok": True, "queued": True}


@api.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user_public(user)


@api.post("/auth/logout")
async def logout(request: Request, response: Response):
    token = request.cookies.get("session_token") or (
        request.headers.get("Authorization", "").replace("Bearer ", "") or None
    )
    if token:
        await db.user_sessions.delete_one({"session_token": token})
    response.delete_cookie("session_token", path="/")
    return {"ok": True}


# ---------- Midtrans payment ----------
MIDTRANS_SERVER_KEY = os.environ.get("MIDTRANS_SERVER_KEY", "")
MIDTRANS_CLIENT_KEY = os.environ.get("MIDTRANS_CLIENT_KEY", "")
MIDTRANS_IS_PRODUCTION = os.environ.get("MIDTRANS_IS_PRODUCTION", "false").lower() == "true"
MIDTRANS_SNAP_URL = ("https://app.midtrans.com" if MIDTRANS_IS_PRODUCTION else "https://app.sandbox.midtrans.com") + "/snap/v1/transactions"
MIDTRANS_STATUS_URL = ("https://api.midtrans.com" if MIDTRANS_IS_PRODUCTION else "https://api.sandbox.midtrans.com") + "/v2"
PREMIUM_PLANS = {
    "monthly": {"days": 30, "amount": 149000, "label": "Premium Bulanan"},
    "yearly": {"days": 365, "amount": 1290000, "label": "Premium Tahunan"},
}


class CheckoutIn(BaseModel):
    plan: str


@api.post("/subscription/checkout")
async def create_checkout(body: CheckoutIn, user: dict = Depends(get_current_user)):
    plan = PREMIUM_PLANS.get(body.plan)
    if not plan:
        raise HTTPException(status_code=400, detail="Paket tidak valid")
    if not MIDTRANS_SERVER_KEY:
        raise HTTPException(status_code=500, detail="Pembayaran belum dikonfigurasi")
    # Resolve price from admin settings (supports promo pricing + promo expiry)
    settings = await _get_settings()
    promo_active = _promo_active(settings)
    if body.plan == "monthly":
        amount = _effective_price(settings.get("monthlyPrice", plan["amount"]),
                                  settings.get("monthlyPromo", 0), promo_active)
    else:
        amount = _effective_price(settings.get("yearlyPrice", plan["amount"]),
                                  settings.get("yearlyPromo", 0), promo_active)
    if amount <= 0:
        amount = plan["amount"]
    order_id = f"PF-{user['user_id'][:12]}-{uuid.uuid4().hex[:10]}"
    await db.orders.insert_one({
        "order_id": order_id, "user_id": user["user_id"], "plan": body.plan,
        "plan_days": plan["days"], "gross_amount": amount, "status": "pending",
        "premium_until": None, "created_at": now_iso(), "last_notification": None,
    })
    payload = {
        "transaction_details": {"order_id": order_id, "gross_amount": amount},
        "enabled_payments": ["other_qris", "gopay", "bank_transfer"],
        "item_details": [{"id": body.plan, "price": amount, "quantity": 1, "name": plan["label"]}],
        "customer_details": {"first_name": user.get("name") or "Pengguna", "email": user.get("email")},
        "expiry": {"unit": "day", "duration": 1},
    }
    auth = base64.b64encode(f"{MIDTRANS_SERVER_KEY}:".encode()).decode()
    try:
        async with httpx.AsyncClient(timeout=20) as client:
            r = await client.post(MIDTRANS_SNAP_URL, json=payload, headers={
                "Accept": "application/json", "Content-Type": "application/json",
                "Authorization": f"Basic {auth}",
            })
    except Exception:
        await db.orders.update_one({"order_id": order_id}, {"$set": {"status": "create_failed"}})
        raise HTTPException(status_code=502, detail="Gagal menghubungi Midtrans")
    if r.status_code >= 400:
        await db.orders.update_one({"order_id": order_id}, {"$set": {"status": "create_failed"}})
        logger.error("Midtrans snap error %s: %s", r.status_code, r.text)
        raise HTTPException(status_code=502, detail="Gagal membuat transaksi pembayaran")
    data = r.json()
    await db.orders.update_one({"order_id": order_id}, {"$set": {"snap_token": data["token"], "redirect_url": data.get("redirect_url")}})
    return {"order_id": order_id, "token": data["token"], "redirect_url": data.get("redirect_url"), "client_key": MIDTRANS_CLIENT_KEY, "production": MIDTRANS_IS_PRODUCTION}


async def _midtrans_status(order_id: str) -> dict:
    auth = base64.b64encode(f"{MIDTRANS_SERVER_KEY}:".encode()).decode()
    async with httpx.AsyncClient(timeout=15) as client:
        r = await client.get(f"{MIDTRANS_STATUS_URL}/{order_id}/status", headers={"Authorization": f"Basic {auth}", "Accept": "application/json"})
    return r.json() if r.status_code < 400 else {}


async def _activate_premium(order: dict) -> str:
    now = datetime.now(timezone.utc)
    user = await db.users.find_one({"user_id": order["user_id"]}, {"_id": 0})
    base = now
    if user and user.get("subscriptionExpiry"):
        try:
            cur = datetime.fromisoformat(user["subscriptionExpiry"])
            if cur.tzinfo is None:
                cur = cur.replace(tzinfo=timezone.utc)
            if cur > now:
                base = cur
        except Exception:
            pass
    new_expiry = (base + timedelta(days=order["plan_days"])).isoformat()
    await db.users.update_one({"user_id": order["user_id"]}, {"$set": {"subscriptionTier": "premium", "subscriptionExpiry": new_expiry}})
    return new_expiry


@api.post("/midtrans/notification")
async def midtrans_notification(request: Request):
    body = await request.json()
    order_id = str(body.get("order_id", ""))
    status_code = str(body.get("status_code", ""))
    gross = str(body.get("gross_amount", ""))
    signature = str(body.get("signature_key", ""))
    expected = hashlib.sha512((order_id + status_code + gross + MIDTRANS_SERVER_KEY).encode()).hexdigest()
    if not order_id or not hmac.compare_digest(expected, signature):
        raise HTTPException(status_code=401, detail="invalid signature")
    order = await db.orders.find_one({"order_id": order_id}, {"_id": 0})
    if not order:
        raise HTTPException(status_code=404, detail="unknown order")
    if gross != f'{order["gross_amount"]:.2f}':
        raise HTTPException(status_code=400, detail="amount mismatch")
    verified = await _midtrans_status(order_id)
    tstatus = str(verified.get("transaction_status", "")).lower()
    fraud = str(verified.get("fraud_status", "")).lower()
    paid = str(verified.get("status_code")) == "200" and tstatus in {"settlement", "capture"} and fraud in {"accept", ""}
    if paid:
        res = await db.orders.update_one({"order_id": order_id, "status": {"$ne": "paid"}}, {"$set": {"status": "paid", "paid_at": now_iso(), "last_notification": body}})
        if res.modified_count:
            expiry = await _activate_premium(order)
            await db.orders.update_one({"order_id": order_id}, {"$set": {"premium_until": expiry}})
    elif tstatus in {"deny", "cancel", "expire"}:
        await db.orders.update_one({"order_id": order_id, "status": {"$ne": "paid"}}, {"$set": {"status": tstatus, "last_notification": body}})
    else:
        await db.orders.update_one({"order_id": order_id}, {"$set": {"status": tstatus or "pending", "last_notification": body}})
    return {"ok": True}


@api.get("/subscription/order/{order_id}")
async def get_order(order_id: str, user: dict = Depends(get_current_user)):
    order = await db.orders.find_one({"order_id": order_id, "user_id": user["user_id"]}, {"_id": 0})
    if not order:
        raise HTTPException(status_code=404, detail="Order tidak ditemukan")
    return {"order_id": order_id, "status": order["status"], "plan": order.get("plan"), "premium_until": order.get("premium_until")}


@api.get("/subscription/orders")
async def list_orders(user: dict = Depends(get_current_user)):
    orders = await db.orders.find(
        {"user_id": user["user_id"]}, {"_id": 0, "snap_token": 0, "last_notification": 0, "redirect_url": 0}
    ).sort("created_at", -1).to_list(500)
    return orders


# ---------- Admin settings ----------
SETTINGS_ID = "app_settings"
DEFAULT_SETTINGS = {
    "appName": "ProFinance Interior",
    "supportEmail": "",
    "supportWhatsapp": "",
    "announcement": "",
    "maintenanceMode": False,
    "monthlyPrice": 149000,
    "monthlyPromo": 149000,
    "yearlyPrice": 1290000,
    "yearlyPromo": 1290000,
    "promoActive": False,
    "promoEndsAt": "",
    "announcementTheme": "info",
}


class SettingsIn(BaseModel):
    appName: Optional[str] = None
    supportEmail: Optional[str] = None
    supportWhatsapp: Optional[str] = None
    announcement: Optional[str] = None
    announcementTheme: Optional[Literal["info", "promo", "warning"]] = None
    maintenanceMode: Optional[bool] = None
    monthlyPrice: Optional[int] = None
    monthlyPromo: Optional[int] = None
    yearlyPrice: Optional[int] = None
    yearlyPromo: Optional[int] = None
    promoActive: Optional[bool] = None
    promoEndsAt: Optional[str] = None


async def require_admin(user: dict = Depends(get_current_user)):
    if (user.get("email") or "").strip().lower() not in OWNER_EMAILS:
        raise HTTPException(status_code=403, detail="Akses khusus admin")
    return user


async def _get_settings() -> dict:
    doc = await db.settings.find_one({"id": SETTINGS_ID}, {"_id": 0})
    if not doc:
        doc = {"id": SETTINGS_ID, **DEFAULT_SETTINGS}
        await db.settings.insert_one(dict(doc))
    return {**DEFAULT_SETTINGS, **{k: v for k, v in doc.items() if k in DEFAULT_SETTINGS}}


# ---------- Work Categories (admin-managed) ----------
DEFAULT_WORK_CATEGORIES = ["Residensial", "Komersial", "Kantor"]


class WorkCategoryIn(BaseModel):
    name: str = Field(..., min_length=1, max_length=50)
    imageUrl: Optional[str] = None


async def _seed_work_categories():
    if await db.work_categories.count_documents({}) == 0:
        await db.work_categories.insert_many([
            {
                "id": str(uuid.uuid4()), "name": n, "imageUrl": None,
                "createdAt": now_iso(), "updatedAt": now_iso(),
            }
            for n in DEFAULT_WORK_CATEGORIES
        ])
        logger.info("Seeded default work categories")


@api.get("/work-categories")
async def list_work_categories():
    """Public: list of project work categories (with optional images)."""
    docs = await db.work_categories.find({}, {"_id": 0}).sort("createdAt", 1).to_list(500)
    if not docs:
        await _seed_work_categories()
        docs = await db.work_categories.find({}, {"_id": 0}).sort("createdAt", 1).to_list(500)
    return docs


@api.post("/admin/work-categories")
async def create_work_category(body: WorkCategoryIn, user: dict = Depends(require_admin)):
    name = body.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Nama kategori wajib diisi")
    if await db.work_categories.find_one({"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}}):
        raise HTTPException(status_code=400, detail="Kategori dengan nama ini sudah ada")
    doc = {
        "id": str(uuid.uuid4()), "name": name,
        "imageUrl": body.imageUrl or None,
        "createdAt": now_iso(), "updatedAt": now_iso(),
    }
    await db.work_categories.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api.put("/admin/work-categories/{cat_id}")
async def update_work_category(cat_id: str, body: WorkCategoryIn, user: dict = Depends(require_admin)):
    name = body.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Nama kategori wajib diisi")
    dup = await db.work_categories.find_one(
        {"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}, "id": {"$ne": cat_id}}
    )
    if dup:
        raise HTTPException(status_code=400, detail="Kategori dengan nama ini sudah ada")
    res = await db.work_categories.update_one(
        {"id": cat_id},
        {"$set": {"name": name, "imageUrl": body.imageUrl or None, "updatedAt": now_iso()}},
    )
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Kategori tidak ditemukan")
    return await db.work_categories.find_one({"id": cat_id}, {"_id": 0})


@api.delete("/admin/work-categories/{cat_id}")
async def delete_work_category(cat_id: str, user: dict = Depends(require_admin)):
    res = await db.work_categories.delete_one({"id": cat_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Kategori tidak ditemukan")
    return {"ok": True}


@api.get("/admin/settings")
async def get_admin_settings(user: dict = Depends(require_admin)):
    return await _get_settings()


@api.put("/admin/settings")
async def update_admin_settings(body: SettingsIn, user: dict = Depends(require_admin)):
    updates = {k: v for k, v in body.model_dump().items() if v is not None}
    if updates:
        await db.settings.update_one({"id": SETTINGS_ID}, {"$set": updates}, upsert=True)
    return await _get_settings()


@api.get("/settings/public")
async def public_settings():
    s = await _get_settings()
    keys = (
        "appName", "announcement", "announcementTheme", "maintenanceMode", "supportWhatsapp", "supportEmail",
        "monthlyPrice", "monthlyPromo", "yearlyPrice", "yearlyPromo", "promoActive", "promoEndsAt",
    )
    return {**{k: s[k] for k in keys}, "serverNow": now_iso()}


def _promo_active(settings: dict) -> bool:
    """Promo counts only when toggled on AND not past its end date (empty end date = no expiry)."""
    if not settings.get("promoActive"):
        return False
    ends = (settings.get("promoEndsAt") or "").strip()
    if not ends:
        return True
    try:
        end_dt = datetime.fromisoformat(ends.replace("Z", "+00:00"))
        if end_dt.tzinfo is None:
            end_dt = end_dt.replace(tzinfo=timezone.utc)
        return datetime.now(timezone.utc) < end_dt
    except Exception:
        return True  # unparseable date -> respect the toggle as-is


def _effective_price(base: int, promo: int, promo_active: bool) -> int:
    """Promo price applies only when active and strictly lower than base (and > 0)."""
    base = int(base or 0)
    promo = int(promo or 0)
    if promo_active and 0 < promo < base:
        return promo
    return base


# ---------- Admin: manajemen pengguna ----------
class GrantPremiumIn(BaseModel):
    days: int = 30  # 0 = permanen


def _sub_status(u: dict, now: datetime) -> tuple[str, Optional[int]]:
    tier = u.get("subscriptionTier", "free")
    if tier != "premium":
        return "free", None
    exp_raw = u.get("subscriptionExpiry")
    exp = None
    if exp_raw:
        try:
            exp = datetime.fromisoformat(exp_raw)
            if exp.tzinfo is None:
                exp = exp.replace(tzinfo=timezone.utc)
        except Exception:
            exp = None
    if exp and exp < now:
        return "expired", None
    days_left = (exp - now).days if exp else None
    return "premium", days_left


@api.get("/admin/users")
async def admin_list_users(q: str = "", page: int = 1, limit: int = 15, user: dict = Depends(require_admin)):
    query = {}
    if q.strip():
        rx = {"$regex": re.escape(q.strip()), "$options": "i"}
        query = {"$or": [{"name": rx}, {"email": rx}, {"phone": rx}]}
    page = max(1, page)
    limit = min(max(1, limit), 50)
    total = await db.users.count_documents(query)
    users = (
        await db.users.find(query, {"_id": 0, "password_hash": 0})
        .sort("created_at", -1)
        .skip((page - 1) * limit)
        .limit(limit)
        .to_list(limit)
    )
    now = datetime.now(timezone.utc)
    items = []
    for u in users:
        uid = u["user_id"]
        project_ids = [p["id"] for p in await db.projects.find({"user_id": uid}, {"_id": 0, "id": 1}).to_list(1000)]
        tx_count = (
            await db.transactions.count_documents({"project_id": {"$in": project_ids}}) if project_ids else 0
        )
        last_session = await db.user_sessions.find_one(
            {"user_id": uid}, {"_id": 0, "created_at": 1}, sort=[("created_at", -1)]
        )
        status, days_left = _sub_status(u, now)
        items.append({
            "user_id": uid,
            "name": u.get("name", ""),
            "email": u.get("email", ""),
            "phone": u.get("phone", ""),
            "authProvider": u.get("authProvider", "email"),
            "subscriptionTier": u.get("subscriptionTier", "free"),
            "subscriptionExpiry": u.get("subscriptionExpiry"),
            "status": status,
            "daysLeft": days_left,
            "isDemo": bool(u.get("isDemo", False)),
            "isOwner": (u.get("email", "") or "").strip().lower() in OWNER_EMAILS,
            "created_at": u.get("created_at"),
            "lastLogin": (last_session or {}).get("created_at"),
            "projectCount": len(project_ids),
            "transactionCount": tx_count,
        })
    pages = max(1, -(-total // limit))
    return {"items": items, "total": total, "page": page, "pages": pages}


@api.post("/admin/users/{target_user_id}/premium")
async def admin_grant_premium(target_user_id: str, body: GrantPremiumIn, user: dict = Depends(require_admin)):
    target = await db.users.find_one({"user_id": target_user_id}, {"_id": 0})
    if not target:
        raise HTTPException(status_code=404, detail="Pengguna tidak ditemukan")
    now = datetime.now(timezone.utc)
    if body.days <= 0:
        expiry = now + timedelta(days=3650)  # permanen
    else:
        base = now
        status, _ = _sub_status(target, now)
        if status == "premium" and target.get("subscriptionExpiry"):
            try:
                cur = datetime.fromisoformat(target["subscriptionExpiry"])
                if cur.tzinfo is None:
                    cur = cur.replace(tzinfo=timezone.utc)
                base = max(now, cur)
            except Exception:
                pass
        expiry = base + timedelta(days=body.days)
    await db.users.update_one(
        {"user_id": target_user_id},
        {"$set": {"subscriptionTier": "premium", "subscriptionExpiry": expiry.isoformat()}},
    )
    await db.notifications.insert_one({
        "id": str(uuid.uuid4()), "user_id": target_user_id, "type": "admin",
        "title": "Akses Premium diaktifkan admin",
        "body": "Selamat! Akun Anda kini Premium." + ("" if body.days <= 0 else f" Berlaku hingga {expiry.date().isoformat()}."),
        "read": False, "createdAt": now_iso(),
    })
    logger.info("Admin %s granted premium (%s days) to %s", user.get("email"), body.days, target.get("email"))
    return {"ok": True, "subscriptionExpiry": expiry.isoformat()}


@api.post("/admin/users/{target_user_id}/revoke")
async def admin_revoke_premium(target_user_id: str, user: dict = Depends(require_admin)):
    target = await db.users.find_one({"user_id": target_user_id}, {"_id": 0})
    if not target:
        raise HTTPException(status_code=404, detail="Pengguna tidak ditemukan")
    if (target.get("email", "") or "").strip().lower() in OWNER_EMAILS:
        raise HTTPException(status_code=400, detail="Akun owner tidak dapat dicabut Premium-nya")
    await db.users.update_one(
        {"user_id": target_user_id},
        {"$set": {"subscriptionTier": "free", "subscriptionExpiry": None}},
    )
    await db.notifications.insert_one({
        "id": str(uuid.uuid4()), "user_id": target_user_id, "type": "admin",
        "title": "Akses Premium berakhir",
        "body": "Status langganan Anda dikembalikan ke Free oleh admin.",
        "read": False, "createdAt": now_iso(),
    })
    logger.info("Admin %s revoked premium from %s", user.get("email"), target.get("email"))
    return {"ok": True}


@api.get("/notifications")
async def list_notifications(user: dict = Depends(get_current_user)):
    notifs = await db.notifications.find(
        {"user_id": user["user_id"]}, {"_id": 0}
    ).sort("createdAt", -1).to_list(100)
    return notifs


@api.post("/notifications/{notif_id}/read")
async def mark_notification_read(notif_id: str, user: dict = Depends(get_current_user)):
    await db.notifications.update_one({"id": notif_id, "user_id": user["user_id"]}, {"$set": {"read": True}})
    return {"ok": True}


async def _renewal_reminders_job():
    now = datetime.now(timezone.utc)
    soon = now + timedelta(days=7)
    users = await db.users.find({"subscriptionTier": "premium"}, {"_id": 0}).to_list(10000)
    created = 0
    for u in users:
        if u.get("email", "").strip().lower() in OWNER_EMAILS or u.get("isDemo"):
            continue
        exp_raw = u.get("subscriptionExpiry")
        if not exp_raw:
            continue
        try:
            exp = datetime.fromisoformat(exp_raw)
            if exp.tzinfo is None:
                exp = exp.replace(tzinfo=timezone.utc)
        except Exception:
            continue
        if exp > soon:
            continue  # not near expiry yet
        days_left = (exp - now).days
        expiry_key = exp.date().isoformat()
        exists = await db.notifications.find_one({"user_id": u["user_id"], "type": "renewal", "expiryKey": expiry_key})
        if exists:
            continue
        if days_left < 0:
            title = "Premium Anda telah berakhir"
            body = "Perpanjang sekarang untuk kembali mengakses Progress, Portal Klien, dan Laporan PDF."
        else:
            title = f"Premium berakhir dalam {days_left} hari"
            body = "Perpanjang langganan agar akses Premium Anda tidak terputus."
        await db.notifications.insert_one({
            "id": str(uuid.uuid4()), "user_id": u["user_id"], "type": "renewal",
            "title": title, "body": body, "daysLeft": days_left, "expiryKey": expiry_key,
            "read": False, "createdAt": now_iso(),
        })
        created += 1
    logger.info("Renewal reminders created: %s", created)


@api.post("/cron/renewal-reminders")
async def cron_renewal_reminders(request: Request, background: BackgroundTasks):
    # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
    secret = os.environ.get("WEBHOOK_CRON_SECRET")
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else ""
    if not secret or not hmac.compare_digest(token, secret):
        raise HTTPException(status_code=401, detail="Unauthorized")
    background.add_task(_renewal_reminders_job)
    return {"ok": True, "queued": True}


# ---------- Subscription (mockup) ----------


# ---------- Projects ----------
@api.post("/projects")
async def create_project(body: ProjectIn, user: dict = Depends(get_current_user)):
    doc = {
        "id": str(uuid.uuid4()),
        "user_id": user["user_id"],
        **body.model_dump(),
        "portalSlug": f"{slugify(body.name)}-{secrets.token_hex(3)}",
        "createdAt": now_iso(),
    }
    await db.projects.insert_one(doc)
    doc.pop("_id", None)
    summary = await compute_summary(doc)
    return {**{k: v for k, v in doc.items() if k != "user_id"}, "summary": summary}


@api.get("/projects")
async def list_projects(user: dict = Depends(get_current_user)):
    projects = await db.projects.find({"user_id": user["user_id"]}, {"_id": 0}).sort("createdAt", -1).to_list(1000)
    out = []
    for p in projects:
        summary = await compute_summary(p)
        out.append({**{k: v for k, v in p.items() if k != "user_id"}, "summary": summary})
    return out


@api.get("/dashboard")
async def dashboard(user: dict = Depends(get_current_user)):
    projects = await db.projects.find({"user_id": user["user_id"]}, {"_id": 0}).to_list(1000)
    total_balance = 0
    total_sisa = 0
    total_budget = 0
    total_in = 0
    total_out = 0
    for p in projects:
        s = await compute_summary(p)
        total_balance += s["balance"]
        total_sisa += s["sisaTagihan"]
        total_budget += s["nominal"]
        total_in += s["totalIn"]
        total_out += s["totalOut"]
    # total kasbon tukang aktif
    workers = await db.workers.find(
        {"project_id": {"$in": [p["id"] for p in projects]}}, {"_id": 0}
    ).to_list(5000)
    total_kasbon_aktif = 0
    for w in workers:
        cw = await compute_worker(w)
        if cw["sisaHutang"] > 0:
            total_kasbon_aktif += cw["totalKasbon"]
    return {
        "saldoBersih": total_balance,
        "totalSisaTagihan": total_sisa,
        "totalBudget": total_budget,
        "totalKasbonAktif": total_kasbon_aktif,
        "totalIn": total_in,
        "totalOut": total_out,
        "projectCount": len(projects),
    }


@api.get("/projects/{project_id}")
async def get_project(project_id: str, user: dict = Depends(get_current_user)):
    p = await get_owned_project(project_id, user)
    summary = await compute_summary(p)
    return {**{k: v for k, v in p.items() if k != "user_id"}, "summary": summary}


@api.put("/projects/{project_id}")
async def update_project(project_id: str, body: ProjectIn, user: dict = Depends(get_current_user)):
    await get_owned_project(project_id, user)
    await db.projects.update_one({"id": project_id}, {"$set": body.model_dump()})
    p = await db.projects.find_one({"id": project_id}, {"_id": 0})
    summary = await compute_summary(p)
    return {**{k: v for k, v in p.items() if k != "user_id"}, "summary": summary}


@api.delete("/projects/{project_id}")
async def delete_project(project_id: str, user: dict = Depends(get_current_user)):
    await get_owned_project(project_id, user)
    await db.projects.delete_one({"id": project_id})
    await db.transactions.delete_many({"project_id": project_id})
    await db.workers.delete_many({"project_id": project_id})
    items = await db.work_items.find({"project_id": project_id}, {"_id": 0}).to_list(2000)
    for wi in items:
        await db.progress_entries.delete_many({"workItemId": wi["id"]})
    await db.work_items.delete_many({"project_id": project_id})
    await db.sub_items.delete_many({"project_id": project_id})
    return {"ok": True}


# ---------- Transactions ----------
@api.post("/projects/{project_id}/transactions")
async def add_transaction(project_id: str, body: TransactionIn, user: dict = Depends(get_current_user)):
    await get_owned_project(project_id, user)
    doc = {
        "id": str(uuid.uuid4()),
        "project_id": project_id,
        **body.model_dump(),
        "date": body.date or now_iso(),
        "createdAt": now_iso(),
    }
    await db.transactions.insert_one(doc)
    doc.pop("_id", None)
    return {k: v for k, v in doc.items() if k != "project_id"}


@api.get("/projects/{project_id}/transactions")
async def list_transactions(project_id: str, user: dict = Depends(get_current_user)):
    await get_owned_project(project_id, user)
    txs = await db.transactions.find({"project_id": project_id}, {"_id": 0}).sort("date", -1).to_list(5000)
    return [{k: v for k, v in t.items() if k != "project_id"} for t in txs]


@api.put("/transactions/{tx_id}")
async def update_transaction(tx_id: str, body: TransactionIn, user: dict = Depends(get_current_user)):
    tx = await db.transactions.find_one({"id": tx_id}, {"_id": 0})
    if not tx:
        raise HTTPException(status_code=404, detail="Transaksi tidak ditemukan")
    await get_owned_project(tx["project_id"], user)
    update = body.model_dump()
    if not update.get("date"):
        update["date"] = tx["date"]
    await db.transactions.update_one({"id": tx_id}, {"$set": update})
    doc = await db.transactions.find_one({"id": tx_id}, {"_id": 0})
    return {k: v for k, v in doc.items() if k != "project_id"}


@api.delete("/transactions/{tx_id}")
async def delete_transaction(tx_id: str, user: dict = Depends(get_current_user)):
    tx = await db.transactions.find_one({"id": tx_id}, {"_id": 0})
    if not tx:
        raise HTTPException(status_code=404, detail="Transaksi tidak ditemukan")
    await get_owned_project(tx["project_id"], user)
    await db.transactions.delete_one({"id": tx_id})
    return {"ok": True}


# ---------- Workers ----------
@api.post("/projects/{project_id}/workers")
async def add_worker(project_id: str, body: WorkerIn, user: dict = Depends(get_current_user)):
    await get_owned_project(project_id, user)
    doc = {"id": str(uuid.uuid4()), "project_id": project_id, **body.model_dump(), "createdAt": now_iso()}
    await db.workers.insert_one(doc)
    doc.pop("_id", None)
    return await compute_worker(doc)


@api.get("/projects/{project_id}/workers")
async def list_workers(project_id: str, user: dict = Depends(get_current_user)):
    await get_owned_project(project_id, user)
    workers = await db.workers.find({"project_id": project_id}, {"_id": 0}).to_list(2000)
    return [await compute_worker(w) for w in workers]


@api.put("/workers/{worker_id}")
async def update_worker(worker_id: str, body: WorkerIn, user: dict = Depends(get_current_user)):
    worker = await db.workers.find_one({"id": worker_id}, {"_id": 0})
    if not worker:
        raise HTTPException(status_code=404, detail="Tukang tidak ditemukan")
    await get_owned_project(worker["project_id"], user)
    old_name = worker["name"]
    new_name = body.name
    await db.workers.update_one({"id": worker_id}, {"$set": {"name": new_name, "borongan": body.borongan}})
    if new_name != old_name:
        for prefix in ("Kasbon Tukang", "Pelunasan Tukang"):
            await db.transactions.update_many(
                {"project_id": worker["project_id"], "category": f"{prefix} {old_name}"},
                {"$set": {"category": f"{prefix} {new_name}"}},
            )
    worker = await db.workers.find_one({"id": worker_id}, {"_id": 0})
    return await compute_worker(worker)


@api.post("/workers/{worker_id}/pay")
async def pay_worker(worker_id: str, body: WorkerPayIn, user: dict = Depends(get_current_user)):
    worker = await db.workers.find_one({"id": worker_id}, {"_id": 0})
    if not worker:
        raise HTTPException(status_code=404, detail="Tukang tidak ditemukan")
    await get_owned_project(worker["project_id"], user)
    label = "Kasbon Tukang" if body.type == "kasbon" else "Pelunasan Tukang"
    category = f"{label} {worker['name']}"
    doc = {
        "id": str(uuid.uuid4()),
        "project_id": worker["project_id"],
        "type": "out",
        "amount": body.amount,
        "description": body.description or f"{label} {worker['name']}",
        "date": body.date or now_iso(),
        "category": category,
        "receiptUrl": None,
        "createdAt": now_iso(),
    }
    await db.transactions.insert_one(doc)
    worker = await db.workers.find_one({"id": worker_id}, {"_id": 0})
    return await compute_worker(worker)


@api.delete("/workers/{worker_id}")
async def delete_worker(worker_id: str, user: dict = Depends(get_current_user)):
    worker = await db.workers.find_one({"id": worker_id}, {"_id": 0})
    if not worker:
        raise HTTPException(status_code=404, detail="Tukang tidak ditemukan")
    await get_owned_project(worker["project_id"], user)
    await db.workers.delete_one({"id": worker_id})
    return {"ok": True}


# ---------- Work Items (Premium) ----------
@api.post("/projects/{project_id}/workitems")
async def add_workitem(project_id: str, body: WorkItemIn, user: dict = Depends(require_premium)):
    await get_owned_project(project_id, user)
    doc = {"id": str(uuid.uuid4()), "project_id": project_id, **body.model_dump(), "createdAt": now_iso()}
    await db.work_items.insert_one(doc)
    p = await db.projects.find_one({"id": project_id}, {"_id": 0})
    items = await compute_workitems(p)
    return next(i for i in items if i["id"] == doc["id"])


@api.get("/projects/{project_id}/workitems")
async def list_workitems(project_id: str, user: dict = Depends(require_premium)):
    p = await get_owned_project(project_id, user)
    return await compute_workitems(p)


@api.put("/workitems/{item_id}")
async def update_workitem(item_id: str, body: WorkItemIn, user: dict = Depends(require_premium)):
    wi = await db.work_items.find_one({"id": item_id}, {"_id": 0})
    if not wi:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    await get_owned_project(wi["project_id"], user)
    await db.work_items.update_one({"id": item_id}, {"$set": body.model_dump()})
    p = await db.projects.find_one({"id": wi["project_id"]}, {"_id": 0})
    items = await compute_workitems(p)
    return next((i for i in items if i["id"] == item_id), {})


@api.delete("/workitems/{item_id}")
async def delete_workitem(item_id: str, user: dict = Depends(require_premium)):
    wi = await db.work_items.find_one({"id": item_id}, {"_id": 0})
    if not wi:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    await get_owned_project(wi["project_id"], user)
    await db.work_items.delete_one({"id": item_id})
    await db.progress_entries.delete_many({"workItemId": item_id})
    await db.sub_items.delete_many({"work_item_id": item_id})
    return {"ok": True}


@api.post("/workitems/{item_id}/subitems")
async def add_subitem(item_id: str, body: SubItemIn, user: dict = Depends(require_premium)):
    wi = await db.work_items.find_one({"id": item_id}, {"_id": 0})
    if not wi:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    await get_owned_project(wi["project_id"], user)
    doc = {
        "id": str(uuid.uuid4()),
        "work_item_id": item_id,
        "project_id": wi["project_id"],
        "name": body.name,
        "harga": body.harga,
        "status": body.status,
        "createdAt": now_iso(),
    }
    await db.sub_items.insert_one(doc)
    doc.pop("_id", None)
    return {k: v for k, v in doc.items() if k != "project_id"}


@api.get("/workitems/{item_id}/subitems")
async def list_subitems(item_id: str, user: dict = Depends(require_premium)):
    wi = await db.work_items.find_one({"id": item_id}, {"_id": 0})
    if not wi:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    await get_owned_project(wi["project_id"], user)
    subs = await db.sub_items.find({"work_item_id": item_id}, {"_id": 0}).sort("createdAt", 1).to_list(2000)
    return [{k: v for k, v in s.items() if k != "project_id"} for s in subs]


@api.put("/subitems/{sub_id}")
async def update_subitem(sub_id: str, body: SubItemIn, user: dict = Depends(require_premium)):
    sub = await db.sub_items.find_one({"id": sub_id}, {"_id": 0})
    if not sub:
        raise HTTPException(status_code=404, detail="Sub item tidak ditemukan")
    await get_owned_project(sub["project_id"], user)
    await db.sub_items.update_one(
        {"id": sub_id},
        {"$set": {"name": body.name, "harga": body.harga, "status": body.status}},
    )
    doc = await db.sub_items.find_one({"id": sub_id}, {"_id": 0})
    return {k: v for k, v in doc.items() if k != "project_id"}


@api.delete("/subitems/{sub_id}")
async def delete_subitem(sub_id: str, user: dict = Depends(require_premium)):
    sub = await db.sub_items.find_one({"id": sub_id}, {"_id": 0})
    if not sub:
        raise HTTPException(status_code=404, detail="Sub item tidak ditemukan")
    await get_owned_project(sub["project_id"], user)
    await db.sub_items.delete_one({"id": sub_id})
    return {"ok": True}


@api.post("/workitems/{item_id}/progress")
async def add_progress(item_id: str, body: ProgressIn, user: dict = Depends(require_premium)):
    wi = await db.work_items.find_one({"id": item_id}, {"_id": 0})
    if not wi:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    await get_owned_project(wi["project_id"], user)
    doc = {
        "id": str(uuid.uuid4()),
        "workItemId": item_id,
        "subItemId": body.subItemId,
        "date": body.date or now_iso(),
        "progress": max(0, min(100, body.progress)),
        "notes": body.notes,
        "photoUrls": body.photoUrls,
        "createdAt": now_iso(),
    }
    await db.progress_entries.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api.get("/workitems/{item_id}/progress")
async def list_progress(item_id: str, subItemId: Optional[str] = Query(None), user: dict = Depends(require_premium)):
    wi = await db.work_items.find_one({"id": item_id}, {"_id": 0})
    if not wi:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    await get_owned_project(wi["project_id"], user)
    q = {"workItemId": item_id}
    if subItemId:
        q["subItemId"] = subItemId
    else:
        q["$or"] = [{"subItemId": None}, {"subItemId": {"$exists": False}}]
    entries = await db.progress_entries.find(q, {"_id": 0}).sort("date", 1).to_list(2000)
    return entries


@api.put("/progress/{entry_id}")
async def update_progress(entry_id: str, body: ProgressIn, user: dict = Depends(require_premium)):
    entry = await db.progress_entries.find_one({"id": entry_id}, {"_id": 0})
    if not entry:
        raise HTTPException(status_code=404, detail="Entri tidak ditemukan")
    wi = await db.work_items.find_one({"id": entry["workItemId"]}, {"_id": 0})
    if not wi:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    await get_owned_project(wi["project_id"], user)
    await db.progress_entries.update_one({"id": entry_id}, {"$set": {
        "progress": max(0, min(100, body.progress)),
        "notes": body.notes,
        "date": body.date or entry.get("date"),
        "photoUrls": body.photoUrls,
    }})
    return await db.progress_entries.find_one({"id": entry_id}, {"_id": 0})


@api.delete("/progress/{entry_id}")
async def delete_progress(entry_id: str, user: dict = Depends(require_premium)):
    entry = await db.progress_entries.find_one({"id": entry_id}, {"_id": 0})
    if not entry:
        raise HTTPException(status_code=404, detail="Entri tidak ditemukan")
    wi = await db.work_items.find_one({"id": entry["workItemId"]}, {"_id": 0})
    if wi:
        await get_owned_project(wi["project_id"], user)
    await db.progress_entries.delete_one({"id": entry_id})
    return {"ok": True}


@api.get("/projects/{project_id}/progress-summary")
async def progress_summary(project_id: str, user: dict = Depends(require_premium)):
    p = await get_owned_project(project_id, user)
    return await _compute_scurve(p)


async def _compute_scurve(p: dict):
    project_id = p["id"]
    data = await compute_progress(p)
    items = data["items"]
    all_entries = await db.progress_entries.find(
        {"workItemId": {"$in": [i["id"] for i in items]}}, {"_id": 0}
    ).sort("date", 1).to_list(5000)

    def parse(d):
        try:
            return datetime.fromisoformat(str(d)[:19])
        except Exception:
            return None

    proj_start = parse(p.get("tanggalMulai"))
    proj_end = parse(p.get("targetSelesai"))

    leaves = []  # (workItemId, subItemId, weight, start_dt, end_dt)
    for i in items:
        s = parse(i.get("startDate")) or proj_start
        e = parse(i.get("endDate")) or proj_end
        if i["hasSubs"]:
            for su in i["subItems"]:
                leaves.append((i["id"], su["id"], su["weight"], s, e))
        else:
            leaves.append((i["id"], None, i["weight"], s, e))

    entry_dts = [parse(e["date"]) for e in all_entries if parse(e["date"])]
    starts = [s for (_, _, _, s, _) in leaves if s]
    ends = [e for (_, _, _, _, e) in leaves if e]
    candidates = starts + ends + entry_dts
    if candidates:
        start_dt = min(starts + entry_dts) if (starts + entry_dts) else min(candidates)
        end_dt = max(ends + entry_dts) if (ends + entry_dts) else max(candidates)
    else:
        start_dt = end_dt = datetime.now().replace(microsecond=0)
    if end_dt < start_dt:
        end_dt = start_dt
    span_days = max((end_dt - start_dt).days, 1)

    def planned_at(dt):
        tot = 0.0
        for (_, _, w, s, e) in leaves:
            ls = s or start_dt
            le = e or end_dt
            if le <= ls:
                frac = 1.0 if dt >= ls else 0.0
            else:
                frac = max(0.0, min(1.0, (dt - ls).days / (le - ls).days))
            tot += w * frac
        return tot

    wmap = {(wi, si): w for (wi, si, w, _, _) in leaves}

    def actual_at(dstr):
        latest = {}
        for e in all_entries:
            if e["date"][:10] <= dstr:
                latest[(e["workItemId"], e.get("subItemId") or None)] = e["progress"]
        return sum(wmap.get(k, 0) * prog / 100 for k, prog in latest.items())

    ticks = set()
    for k in range(11):
        ticks.add((start_dt + timedelta(days=round(span_days * k / 10))).date().isoformat())
    for e in all_entries:
        pd = parse(e["date"])
        if pd:
            ticks.add(pd.date().isoformat())

    curve = []
    for d in sorted(ticks):
        dt = datetime.fromisoformat(d)
        curve.append({
            "date": d,
            "planned": round(min(100.0, planned_at(dt)), 1),
            "actual": round(min(100.0, actual_at(d)), 1),
        })

    today = datetime.now().replace(microsecond=0)
    return {
        "totalProgress": data["totalProgress"],
        "plannedProgress": round(min(100.0, planned_at(today)), 1),
        "items": items,
        "curve": curve,
        "rab": data["rab"],
        "rabTotal": data["rabTotal"],
        "totalItemValue": data["totalItemValue"],
        "completedValue": data["completedValue"],
        "rabBaseline": data["rabBaseline"],
        "start": start_dt.date().isoformat(),
        "end": end_dt.date().isoformat(),
    }


@api.put("/projects/{project_id}/rab-total")
async def set_rab_total(project_id: str, body: RabIn, user: dict = Depends(get_current_user)):
    await get_owned_project(project_id, user)
    await db.projects.update_one({"id": project_id}, {"$set": {"rabTotal": body.rabTotal}})
    return {"rabTotal": body.rabTotal}


@api.post("/projects/{project_id}/rab-baseline")
async def save_rab_baseline(project_id: str, user: dict = Depends(require_premium)):
    p = await get_owned_project(project_id, user)
    data = await compute_progress(p)
    baseline = {
        "savedAt": now_iso(),
        "rabTotal": data["rabTotal"],
        "totalItemValue": data["totalItemValue"],
        "items": [{"id": i["id"], "name": i["name"], "value": i["nilai"]} for i in data["items"]],
    }
    await db.projects.update_one({"id": project_id}, {"$set": {"rabBaseline": baseline}})
    return baseline


@api.delete("/projects/{project_id}/rab-baseline")
async def delete_rab_baseline(project_id: str, user: dict = Depends(require_premium)):
    await get_owned_project(project_id, user)
    await db.projects.update_one({"id": project_id}, {"$unset": {"rabBaseline": ""}})
    return {"ok": True}


@api.get("/projects/{project_id}/portal-link")
async def portal_link(project_id: str, user: dict = Depends(get_current_user)):
    p = await get_owned_project(project_id, user)
    slug = await ensure_portal_slug(p)
    return {"slug": slug}


@api.post("/projects/{project_id}/portal-link/reset")
async def reset_portal_link(project_id: str, user: dict = Depends(get_current_user)):
    p = await get_owned_project(project_id, user)
    slug = f"{slugify(p.get('name'))}-{secrets.token_hex(3)}"
    await db.projects.update_one({"id": project_id}, {"$set": {"portalSlug": slug}})
    return {"slug": slug}


# ---------- RAB Builder / Surat Penawaran (Premium) ----------
@api.get("/rab/company-template")
async def get_company_template(user: dict = Depends(require_premium)):
    tpl = await db.rab_templates.find_one({"user_id": user["user_id"]}, {"_id": 0, "user_id": 0})
    return tpl or {}


@api.put("/rab/company-template")
async def save_company_template(body: CompanyTemplateIn, user: dict = Depends(require_premium)):
    data = body.model_dump()
    # Guard against oversized signature payloads (base64 data URI)
    if data.get("signatureImage") and len(data["signatureImage"]) > 3_000_000:
        raise HTTPException(status_code=400, detail="Gambar tanda tangan terlalu besar (maks ~2MB)")
    data["user_id"] = user["user_id"]
    data["updatedAt"] = now_iso()
    await db.rab_templates.update_one({"user_id": user["user_id"]}, {"$set": data}, upsert=True)
    return {"ok": True, **{k: v for k, v in data.items() if k not in ("user_id",)}}


@api.post("/rab")
async def create_rab(body: RabCreateIn, user: dict = Depends(require_premium)):
    rab = body.rab.model_dump()
    if rab.get("signatureImage") and len(rab["signatureImage"]) > 3_000_000:
        raise HTTPException(status_code=400, detail="Gambar tanda tangan terlalu besar (maks ~2MB)")
    comp = compute_rab(rab)
    pid = str(uuid.uuid4())
    doc = {
        "id": pid, "user_id": user["user_id"], "name": body.projectName,
        "owner": rab.get("clientName", ""), "companyName": rab.get("companyName", ""),
        "alamatProyek": rab.get("clientAddress", ""), "category": body.category,
        "status": "Prospek", "nominal": comp["grandTotal"], "rabTotal": comp["totalItems"],
        "rab": rab,
        "portalSlug": f"{slugify(body.projectName)}-{secrets.token_hex(3)}",
        "rabSlug": f"{slugify(body.projectName)}-rab-{secrets.token_hex(3)}",
        "tanggalMulai": now_iso(), "targetSelesai": None, "thumbnail": None,
        "createdAt": now_iso(),
    }
    await db.projects.insert_one(doc)
    doc.pop("_id", None)
    return {**{k: v for k, v in doc.items() if k != "user_id"}, "computed": comp}


@api.get("/projects/{project_id}/rab")
async def get_rab(project_id: str, user: dict = Depends(require_premium)):
    p = await get_owned_project(project_id, user)
    rab = p.get("rab") or {}
    return {
        "projectId": p["id"], "projectName": p.get("name"), "category": p.get("category"),
        "status": p.get("status"), "rab": rab, "computed": compute_rab(rab),
        "rabSlug": p.get("rabSlug"),
    }


@api.put("/projects/{project_id}/rab")
async def update_rab(project_id: str, body: RabCreateIn, user: dict = Depends(require_premium)):
    p = await get_owned_project(project_id, user)
    rab = body.rab.model_dump()
    if rab.get("signatureImage") and len(rab["signatureImage"]) > 3_000_000:
        raise HTTPException(status_code=400, detail="Gambar tanda tangan terlalu besar (maks ~2MB)")
    comp = compute_rab(rab)
    updates = {
        "rab": rab, "name": body.projectName, "category": body.category,
        "owner": rab.get("clientName", ""), "companyName": rab.get("companyName", ""),
        "alamatProyek": rab.get("clientAddress", ""),
    }
    if p.get("status") == "Prospek":
        updates["nominal"] = comp["grandTotal"]
        updates["rabTotal"] = comp["totalItems"]
    await db.projects.update_one({"id": project_id}, {"$set": updates})
    return {"ok": True, "computed": comp, "rabSlug": p.get("rabSlug")}


@api.post("/projects/{project_id}/deal")
async def deal_project(project_id: str, user: dict = Depends(require_premium)):
    p = await get_owned_project(project_id, user)
    rab = p.get("rab") or {}
    comp = compute_rab(rab)
    # Populate Progress (work_items + sub_items) once, so RAB values flow into progress tracking.
    existing = await db.work_items.count_documents({"project_id": project_id})
    if existing == 0:
        for sec in comp["sections"]:
            wi_id = str(uuid.uuid4())
            await db.work_items.insert_one({
                "id": wi_id, "project_id": project_id, "name": sec.get("name") or "Pekerjaan",
                "nilai": sec.get("subtotal", 0), "startDate": None, "endDate": None,
                "createdAt": now_iso(),
            })
            for si in sec.get("subItems", []):
                await db.sub_items.insert_one({
                    "id": str(uuid.uuid4()), "work_item_id": wi_id, "project_id": project_id,
                    "name": si.get("name") or "Item", "harga": si.get("nilai", 0),
                    "status": False, "createdAt": now_iso(),
                })
    await db.projects.update_one({"id": project_id}, {"$set": {
        "status": "Berjalan", "rabTotal": comp["totalItems"], "nominal": comp["grandTotal"],
        "dealAt": now_iso(),
    }})
    p = await db.projects.find_one({"id": project_id}, {"_id": 0})
    summary = await compute_summary(p)
    return {**{k: v for k, v in p.items() if k != "user_id"}, "summary": summary}


@api.get("/projects/{project_id}/rab/share-link")
async def rab_share_link(project_id: str, user: dict = Depends(require_premium)):
    p = await get_owned_project(project_id, user)
    slug = await ensure_rab_slug(p)
    return {"slug": slug}


@api.get("/projects/{project_id}/rab/pdf")
async def rab_pdf(project_id: str, request: Request, auth: Optional[str] = Query(None)):
    user = await _user_from_request_or_query(request, auth)
    await require_premium(user)
    p = await get_owned_project(project_id, user)
    rab = p.get("rab") or {}
    pdf_bytes = build_rab_pdf(p, rab, compute_rab(rab))
    safe = slugify(p.get("name") or "rab")
    return StreamingResponse(
        iter([pdf_bytes]), media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="Penawaran-{safe}.pdf"'},
    )


@api.get("/public/rab/{slug}/pdf")
async def public_rab_pdf(slug: str):
    p = await db.projects.find_one({"rabSlug": slug}, {"_id": 0})
    if not p:
        raise HTTPException(status_code=404, detail="Penawaran tidak ditemukan")
    rab = p.get("rab") or {}
    pdf_bytes = build_rab_pdf(p, rab, compute_rab(rab))
    safe = slugify(p.get("name") or "rab")
    return StreamingResponse(
        iter([pdf_bytes]), media_type="application/pdf",
        headers={"Content-Disposition": f'inline; filename="Penawaran-{safe}.pdf"'},
    )


# ---------- Invoices (Premium) ----------
INVOICE_TYPES = {"proforma", "final", "retention"}


async def _get_owned_invoice(invoice_id: str, user: dict) -> dict:
    inv = await db.invoices.find_one({"id": invoice_id, "user_id": user["user_id"]}, {"_id": 0})
    if not inv:
        raise HTTPException(status_code=404, detail="Invoice tidak ditemukan")
    return inv


@api.get("/projects/{project_id}/invoice-context")
async def invoice_context(project_id: str, user: dict = Depends(require_premium)):
    """Prefill data for the invoice form: client, company template, RAB payment terms,
    payment progress, existing invoices, and a suggested invoice type."""
    p = await get_owned_project(project_id, user)
    summary = await compute_summary(p)
    rab = p.get("rab") or {}
    comp = compute_rab(rab) if rab else {}
    template = await db.rab_templates.find_one({"user_id": user["user_id"]}, {"_id": 0}) or {}

    txs = await db.transactions.find({"project_id": project_id, "type": "in"}, {"_id": 0}).to_list(5000)
    paid = {"Downpayment": 0, "Termin": 0, "Pelunasan": 0}
    for t in txs:
        cat = t.get("category")
        if cat in paid:
            paid[cat] += t.get("amount", 0)

    termins = comp.get("termins", []) or []
    existing = await db.invoices.find({"project_id": project_id}, {"_id": 0}).sort("createdAt", 1).to_list(500)
    proforma_count = sum(1 for e in existing if e.get("type") == "proforma")

    # ---- Suggestion (correct financial workflow) ----
    # DP & progress termins -> Proforma; final settlement -> Invoice (final).
    sisa = summary.get("sisaTagihan", 0)
    suggested_type = "proforma"
    suggested_desc = "Uang Muka (DP)"
    suggested_amount = 0
    if termins:
        # index of the next installment to bill (by count of proforma issued)
        idx = min(proforma_count, len(termins) - 1)
        is_last = idx >= len(termins) - 1
        nxt = termins[idx]
        label = (nxt.get("label") or "").lower()
        suggested_desc = nxt.get("label") or f"Termin {idx + 1}"
        suggested_amount = nxt.get("nominal", 0)
        if is_last or "lunas" in label or "pelunasan" in label:
            suggested_type = "final"
    else:
        # No RAB terms: DP first, then final for the remainder
        if paid["Downpayment"] == 0 and paid["Termin"] == 0 and paid["Pelunasan"] == 0:
            suggested_type = "proforma"
            suggested_desc = "Uang Muka (DP)"
            suggested_amount = round((p.get("nominal", 0) or 0) * 0.5)
        else:
            suggested_type = "final"
            suggested_desc = "Pelunasan"
            suggested_amount = max(0, sisa)

    return {
        "project": {"id": p["id"], "name": p.get("name"), "nominal": p.get("nominal", 0), "status": p.get("status")},
        "client": {
            "clientName": p.get("owner") or rab.get("clientName", ""),
            "clientAddress": p.get("alamatProyek") or rab.get("clientAddress", ""),
            "clientPhone": rab.get("clientPhone", ""),
        },
        "company": {k: template.get(k, "") for k in (
            "companyName", "companyAddress", "companyPhone", "bankName", "bankAccount",
            "bankHolder", "signerLeft", "signerRight", "signatureImage")},
        "ppn": {"ppnEnabled": comp.get("ppnEnabled", False), "ppnPercent": comp.get("ppnPercent", 11)},
        "rabRef": {"quotationNo": rab.get("quotationNo", ""), "quotationDate": rab.get("quotationDate")},
        "termins": termins,
        "summary": {**summary, "paid": paid},
        "suggestion": {"type": suggested_type, "description": suggested_desc, "amount": suggested_amount},
        "existingCount": len(existing),
    }

@api.get("/projects/{project_id}/billing-recap")
async def billing_recap(project_id: str, user: dict = Depends(require_premium)):
    """Billing recap: total billed (invoices sent/paid) vs collected vs retention held."""
    p = await get_owned_project(project_id, user)
    summary = await compute_summary(p)
    invoices = await db.invoices.find(
        {"project_id": project_id, "user_id": user["user_id"]}, {"_id": 0}
    ).to_list(500)
    total_billed = 0       # invoices with status Terkirim/Lunas (actually billed to client)
    draft_amount = 0
    retention_held = 0
    billed_count = 0
    for inv in invoices:
        comp = compute_invoice(inv)
        retention_held += comp.get("retentionAmount", 0)
        if inv.get("status") in ("Terkirim", "Lunas"):
            total_billed += comp.get("amountDue", 0)
            billed_count += 1
        else:
            draft_amount += comp.get("amountDue", 0)
    paid = summary.get("terbayar", 0)
    receivable = max(0, total_billed - paid)
    return {
        "nominal": summary.get("nominal", 0),
        "totalBilled": total_billed,
        "draftAmount": draft_amount,
        "paid": paid,
        "receivable": receivable,
        "retentionHeld": retention_held,
        "sisaTagihan": summary.get("sisaTagihan", 0),
        "invoiceCount": len(invoices),
        "billedCount": billed_count,
    }




def _invoice_public(inv: dict) -> dict:
    doc = {k: v for k, v in inv.items() if k != "user_id"}
    doc["computed"] = compute_invoice(inv)
    return doc


@api.post("/projects/{project_id}/invoices")
async def create_invoice(project_id: str, body: InvoiceIn, user: dict = Depends(require_premium)):
    p = await get_owned_project(project_id, user)
    if body.type not in INVOICE_TYPES:
        raise HTTPException(status_code=400, detail="Tipe invoice tidak valid")
    if not (body.number or "").strip():
        raise HTTPException(status_code=400, detail="Nomor invoice wajib diisi")
    if body.signatureImage and len(body.signatureImage) > 3_000_000:
        raise HTTPException(status_code=400, detail="Gambar tanda tangan terlalu besar (maks ~2MB)")
    doc = {
        "id": str(uuid.uuid4()), "user_id": user["user_id"], "project_id": project_id,
        "projectName": p.get("name"),
        **body.model_dump(),
        "number": body.number.strip(),
        "invoiceDate": body.invoiceDate or now_iso(),
        "createdAt": now_iso(), "updatedAt": now_iso(),
    }
    await db.invoices.insert_one(doc)
    doc.pop("_id", None)
    return _invoice_public(doc)


@api.get("/projects/{project_id}/invoices")
async def list_invoices(project_id: str, user: dict = Depends(require_premium)):
    await get_owned_project(project_id, user)
    docs = await db.invoices.find({"project_id": project_id, "user_id": user["user_id"]}, {"_id": 0}).sort("createdAt", -1).to_list(500)
    return [_invoice_public(d) for d in docs]


@api.get("/invoices/{invoice_id}")
async def get_invoice(invoice_id: str, user: dict = Depends(require_premium)):
    inv = await _get_owned_invoice(invoice_id, user)
    return _invoice_public(inv)


@api.put("/invoices/{invoice_id}")
async def update_invoice(invoice_id: str, body: InvoiceIn, user: dict = Depends(require_premium)):
    await _get_owned_invoice(invoice_id, user)
    if body.type not in INVOICE_TYPES:
        raise HTTPException(status_code=400, detail="Tipe invoice tidak valid")
    if not (body.number or "").strip():
        raise HTTPException(status_code=400, detail="Nomor invoice wajib diisi")
    if body.signatureImage and len(body.signatureImage) > 3_000_000:
        raise HTTPException(status_code=400, detail="Gambar tanda tangan terlalu besar (maks ~2MB)")
    updates = {**body.model_dump(), "number": body.number.strip(), "updatedAt": now_iso()}
    await db.invoices.update_one({"id": invoice_id}, {"$set": updates})
    inv = await db.invoices.find_one({"id": invoice_id}, {"_id": 0})
    return _invoice_public(inv)


@api.delete("/invoices/{invoice_id}")
async def delete_invoice(invoice_id: str, user: dict = Depends(require_premium)):
    await _get_owned_invoice(invoice_id, user)
    await db.invoices.delete_one({"id": invoice_id})
    return {"ok": True}


@api.get("/invoices/{invoice_id}/pdf")
async def invoice_pdf(invoice_id: str, request: Request, auth: Optional[str] = Query(None)):
    user = await _user_from_request_or_query(request, auth)
    await require_premium(user)
    inv = await _get_owned_invoice(invoice_id, user)
    pdf_bytes = build_invoice_pdf(inv, compute_invoice(inv))
    prefix = {"proforma": "Proforma", "final": "Invoice", "retention": "Invoice-Retensi"}.get(inv.get("type"), "Invoice")
    safe = slugify(inv.get("number") or inv.get("projectName") or "invoice")
    return StreamingResponse(
        iter([pdf_bytes]), media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{prefix}-{safe}.pdf"'},
    )




# ---------- RAB Excel import ----------
def _to_int(v):
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return int(v)
    s = re.sub(r"[^0-9]", "", str(v))
    return int(s) if s else 0


@api.get("/rab-template")
async def rab_template(user: dict = Depends(get_current_user)):
    wb = Workbook()
    ws = wb.active
    ws.title = "RAB"
    ws.append(["Item Pekerjaan", "Sub Item", "Harga (Rp)"])
    ws.append(["Pekerjaan Plafon Gypsum", "Rangka Hollow 4x4", 12000000])
    ws.append(["", "Gypsum board 9mm", 8000000])
    ws.append(["", "Finishing compound & cat", 5000000])
    ws.append(["Pekerjaan Lantai", "Keramik granit 60x60", 25000000])
    ws.append(["", "Pemasangan & nat", 6000000])
    ws.append(["Pekerjaan Pengecatan (tanpa rincian)", "", 15000000])
    for i, w in enumerate((34, 30, 16)):
        ws.column_dimensions[chr(65 + i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        iter([buf.read()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Template-RAB.xlsx"'},
    )


@api.post("/projects/{project_id}/workitems/import")
async def import_rab(project_id: str, file: UploadFile = File(...), user: dict = Depends(require_premium)):
    await get_owned_project(project_id, user)
    raw = await file.read()
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="File Excel tidak valid (.xlsx)")
    ws = wb.active
    created_items = 0
    created_subs = 0
    current_id = None
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        if idx == 0 or not row:
            continue
        item_name = str(row[0]).strip() if len(row) > 0 and row[0] not in (None, "") else ""
        sub_name = str(row[1]).strip() if len(row) > 1 and row[1] not in (None, "") else ""
        harga = _to_int(row[2]) if len(row) > 2 else 0
        if not item_name and not sub_name:
            continue
        if item_name:
            current_id = str(uuid.uuid4())
            await db.work_items.insert_one({
                "id": current_id, "project_id": project_id, "name": item_name,
                "nilai": 0 if sub_name else harga, "createdAt": now_iso(),
            })
            created_items += 1
        if sub_name:
            if not current_id:
                current_id = str(uuid.uuid4())
                await db.work_items.insert_one({
                    "id": current_id, "project_id": project_id, "name": "Item Pekerjaan",
                    "nilai": 0, "createdAt": now_iso(),
                })
                created_items += 1
            await db.sub_items.insert_one({
                "id": str(uuid.uuid4()), "work_item_id": current_id, "project_id": project_id,
                "name": sub_name, "harga": harga, "status": False, "createdAt": now_iso(),
            })
            created_subs += 1
    wb.close()
    if created_items == 0 and created_subs == 0:
        raise HTTPException(status_code=400, detail="Tidak ada data valid ditemukan. Gunakan template RAB.")
    return {"items": created_items, "subs": created_subs}


# ---------- Public Client Portal (no auth) ----------
@api.get("/public/portal/{slug}")
async def public_portal(slug: str):
    p = await db.projects.find_one({"portalSlug": slug}, {"_id": 0})
    if not p:
        raise HTTPException(status_code=404, detail="Portal tidak ditemukan")
    summary = await compute_summary(p)
    scurve = await _compute_scurve(p)

    items = []
    name_map = {}
    sub_name = {}
    for i in scurve["items"]:
        name_map[i["id"]] = i["name"]
        subs = []
        for s in i.get("subItems", []):
            sub_name[s["id"]] = s["name"]
            subs.append({"id": s["id"], "name": s["name"], "lastProgress": s["lastProgress"]})
        items.append({
            "id": i["id"], "name": i["name"], "lastProgress": i["lastProgress"],
            "startDate": i["startDate"], "endDate": i["endDate"],
            "hasSubs": i["hasSubs"], "subItems": subs,
        })

    item_ids = [i["id"] for i in scurve["items"]]
    entries = await db.progress_entries.find(
        {"workItemId": {"$in": item_ids}}, {"_id": 0}
    ).sort("date", -1).to_list(3000)
    gallery = []
    for e in entries:
        urls = [u if u.startswith("http") else f"/api/public/portal/{slug}/file?path={quote(u)}" for u in (e.get("photoUrls") or [])]
        if not urls:
            continue
        label = name_map.get(e["workItemId"], "")
        if e.get("subItemId"):
            label = f"{label} › {sub_name.get(e['subItemId'], '')}"
        gallery.append({
            "date": e["date"], "notes": e.get("notes", ""), "label": label,
            "progress": e.get("progress", 0), "photos": urls,
        })

    txs = await db.transactions.find(
        {"project_id": p["id"], "type": "in"}, {"_id": 0}
    ).sort("date", -1).to_list(5000)
    payments = [{
        "date": t["date"], "category": t.get("category", "Lainnya"),
        "description": t.get("description", ""), "amount": t["amount"],
    } for t in txs]

    return {
        "project": {
            "name": p.get("name"), "companyName": p.get("companyName", ""),
            "owner": p.get("owner", ""), "alamatProyek": p.get("alamatProyek", ""),
            "category": p.get("category", ""), "status": p.get("status", ""),
            "targetSelesai": p.get("targetSelesai"), "tanggalMulai": p.get("tanggalMulai"),
            "thumbnail": p.get("thumbnail"),
        },
        "totalProgress": scurve["totalProgress"],
        "plannedProgress": scurve["plannedProgress"],
        "curve": scurve["curve"], "start": scurve["start"], "end": scurve["end"],
        "items": items, "gallery": gallery,
        "payments": payments,
        "terbayar": summary["terbayar"], "sisaTagihan": summary["sisaTagihan"],
        "nominal": summary["nominal"],
    }


@api.get("/public/portal/{slug}/file")
async def public_portal_file(slug: str, path: str = Query(...)):
    p = await db.projects.find_one({"portalSlug": slug}, {"_id": 0})
    if not p:
        raise HTTPException(status_code=404, detail="Portal tidak ditemukan")
    wi = await db.work_items.find({"project_id": p["id"]}, {"_id": 0, "id": 1}).to_list(2000)
    item_ids = [x["id"] for x in wi]
    entries = await db.progress_entries.find(
        {"workItemId": {"$in": item_ids}}, {"_id": 0, "photoUrls": 1}
    ).to_list(5000)
    allowed = set()
    for e in entries:
        for u in (e.get("photoUrls") or []):
            allowed.add(u)
    if path not in allowed:
        raise HTTPException(status_code=404, detail="File tidak ditemukan")
    data, content_type = get_object(path)
    record = await db.files.find_one({"storage_path": path}, {"_id": 0})
    return Response(content=data, media_type=(record or {}).get("content_type", content_type))


@api.get("/public/portal/{slug}/share")
async def portal_share(slug: str, request: Request):
    p = await db.projects.find_one({"portalSlug": slug}, {"_id": 0})
    if not p:
        raise HTTPException(status_code=404, detail="Portal tidak ditemukan")
    host = request.headers.get("x-forwarded-host") or request.headers.get("host") or request.url.hostname
    target_rel = f"/portal/{slug}"
    og_url = f"https://{host}{target_rel}"
    name = html.escape(p.get("name", "Proyek") or "Proyek")
    company = (p.get("companyName") or "").strip()
    company_txt = f" · {company}" if company and company != "-" else ""
    desc = html.escape(f"Pantau progress pekerjaan, kurva-S, dan foto dokumentasi secara realtime{company_txt}.")
    img = html.escape(p.get("thumbnail") or "")
    title = html.escape(f"{p.get('name', 'Proyek')} — Portal Progress Klien")
    img_tags = f'<meta property="og:image" content="{img}"><meta name="twitter:image" content="{img}">' if img else ""
    doc = f"""<!doctype html><html lang="id"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<meta property="og:type" content="website">
<meta property="og:site_name" content="ProFinance Interior">
<meta property="og:title" content="{title}">
<meta property="og:description" content="{desc}">
<meta property="og:url" content="{html.escape(og_url)}">
{img_tags}
<meta name="twitter:card" content="summary_large_image">
<meta name="twitter:title" content="{title}">
<meta name="twitter:description" content="{desc}">
<meta http-equiv="refresh" content="0; url={target_rel}">
<script>window.location.replace({json.dumps(target_rel)});</script>
</head><body style="font-family:sans-serif;background:#0f172a;color:#fff;display:flex;align-items:center;justify-content:center;height:100vh;margin:0">
Mengalihkan ke portal progress…</body></html>"""
    return Response(content=doc, media_type="text/html")


# ---------- Report (Premium) ----------
@api.get("/projects/{project_id}/report")
async def report(project_id: str, user: dict = Depends(require_premium)):
    p = await get_owned_project(project_id, user)
    summary = await compute_summary(p)
    txs = await db.transactions.find({"project_id": project_id}, {"_id": 0}).to_list(5000)
    expense_by_cat = {}
    for t in txs:
        if t["type"] == "out":
            cat = t.get("category", "Lainnya")
            # group tukang categories
            if cat.startswith("Kasbon Tukang") or cat.startswith("Pelunasan Tukang"):
                cat = "Upah & Kasbon Tukang"
            expense_by_cat[cat] = expense_by_cat.get(cat, 0) + t["amount"]
    expense_list = sorted(
        [{"name": k, "value": v} for k, v in expense_by_cat.items()],
        key=lambda x: x["value"], reverse=True,
    )
    return {
        "summary": summary,
        "expenseByCategory": expense_list,
        "inOut": [
            {"name": "Pemasukan", "value": summary["totalIn"]},
            {"name": "Pengeluaran", "value": summary["totalOut"]},
        ],
        "marginBar": [
            {"name": "Nilai Kontrak", "value": summary["nominal"]},
            {"name": "Total Pengeluaran", "value": summary["totalOut"]},
            {"name": "Laba / Margin", "value": summary["balance"]},
        ],
    }


@api.get("/projects/{project_id}/report/pdf")
async def report_pdf(project_id: str, request: Request, auth: Optional[str] = Query(None)):
    # allow token via query param for direct download links
    if auth and not request.headers.get("Authorization"):
        request.scope.setdefault("headers", [])
    user = await _user_from_request_or_query(request, auth)
    await require_premium(user)
    p = await get_owned_project(project_id, user)
    summary = await compute_summary(p)
    txs = await db.transactions.find({"project_id": project_id}, {"_id": 0}).sort("date", 1).to_list(5000)
    workers_raw = await db.workers.find({"project_id": project_id}, {"_id": 0}).to_list(2000)
    workers = [await compute_worker(w) for w in workers_raw]
    _, work_items = await project_total_progress(p)
    entries = await db.progress_entries.find(
        {"workItemId": {"$in": [i["id"] for i in work_items]}}, {"_id": 0}
    ).sort("date", 1).to_list(5000)
    pdf_bytes = build_report_pdf(p, summary, txs, workers, work_items, entries)
    safe_name = "".join(c for c in p.get("name", "laporan") if c.isalnum() or c in " -_")[:40].strip() or "laporan"
    return StreamingResponse(
        iter([pdf_bytes]),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="Laporan-{safe_name}.pdf"'},
    )


@api.get("/projects/{project_id}/progress/pdf")
async def progress_pdf(project_id: str, request: Request, auth: Optional[str] = Query(None)):
    user = await _user_from_request_or_query(request, auth)
    await require_premium(user)
    p = await get_owned_project(project_id, user)
    summary = await compute_summary(p)
    prog = await _compute_scurve(p)
    txs = await db.transactions.find({"project_id": project_id}, {"_id": 0}).sort("date", 1).to_list(5000)
    item_ids = [i["id"] for i in prog.get("items", [])]
    entries = await db.progress_entries.find(
        {"workItemId": {"$in": item_ids}}, {"_id": 0}
    ).sort("date", 1).to_list(5000)
    entries_by_key = {}
    for e in entries:
        sid = e.get("subItemId")
        key = ("sub", sid) if sid else ("item", e.get("workItemId"))
        entries_by_key.setdefault(key, []).append(e)
    pdf_bytes = build_progress_pdf(p, summary, prog, txs, entries_by_key)
    safe_name = "".join(c for c in p.get("name", "progress") if c.isalnum() or c in " -_")[:40].strip() or "progress"
    return StreamingResponse(
        iter([pdf_bytes]),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="Laporan-Progress-{safe_name}.pdf"'},
    )


async def _user_from_request_or_query(request: Request, auth: Optional[str]) -> dict:
    token = request.cookies.get("session_token")
    if not token:
        header = request.headers.get("Authorization")
        if header and header.startswith("Bearer "):
            token = header.split(" ", 1)[1]
    if not token and auth:
        token = auth
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    session = await db.user_sessions.find_one({"session_token": token}, {"_id": 0})
    if not session:
        raise HTTPException(status_code=401, detail="Sesi tidak ditemukan")
    user = await db.users.find_one({"user_id": session["user_id"]}, {"_id": 0})
    if not user:
        raise HTTPException(status_code=401, detail="User tidak ditemukan")
    return user


# ---------- File upload / serve ----------
@api.post("/upload")
async def upload_file(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    ext = (file.filename.split(".")[-1].lower() if "." in file.filename else "bin")
    content_type = MIME_TYPES.get(ext, file.content_type or "application/octet-stream")
    path = f"{APP_NAME}/uploads/{user['user_id']}/{uuid.uuid4()}.{ext}"
    data = await file.read()
    result = put_object(path, data, content_type)
    canonical = result["path"]
    await db.files.insert_one({
        "id": str(uuid.uuid4()),
        "user_id": user["user_id"],
        "storage_path": canonical,
        "original_filename": file.filename,
        "content_type": content_type,
        "size": result.get("size", len(data)),
        "is_deleted": False,
        "created_at": now_iso(),
    })
    return {"path": canonical}


@api.get("/files/{path:path}")
async def serve_file(path: str, request: Request, auth: Optional[str] = Query(None)):
    await _user_from_request_or_query(request, auth)
    record = await db.files.find_one({"storage_path": path, "is_deleted": False}, {"_id": 0})
    if not record:
        raise HTTPException(status_code=404, detail="File tidak ditemukan")
    data, content_type = get_object(path)
    return Response(content=data, media_type=record.get("content_type", content_type))


# ---------- Backup / Export ----------
@api.get("/backup/summary")
async def backup_summary(user: dict = Depends(get_current_user)):
    """Preview counts of what will be included in the backup."""
    data = await backup_mod.gather_user_data(user)
    photo_paths = backup_mod._collect_photo_paths(data)
    return {
        "projects": len(data["projects"]),
        "transactions": len(data["transactions"]),
        "workers": len(data["workers"]),
        "work_items": len(data["work_items"]),
        "progress_entries": len(data["progress_entries"]),
        "photos": len(photo_paths),
    }


@api.get("/backup/export")
async def backup_export(user: dict = Depends(get_current_user)):
    """Generate a full ZIP backup (data + photos) and download to device."""
    data = await backup_mod.gather_user_data(user)
    photo_paths = backup_mod._collect_photo_paths(data)
    zip_path, stats = await run_in_threadpool(backup_mod._assemble_zip, data, user, photo_paths)
    filename = f"profinance-backup-{stats['stamp']}.zip"
    logger.info(f"Backup export for {user.get('email')}: {stats}")

    def _cleanup(p=zip_path):
        try:
            os.unlink(p)
        except Exception:
            pass

    return FileResponse(
        zip_path,
        media_type="application/zip",
        filename=filename,
        background=BackgroundTask(_cleanup),
    )


@api.get("/projects/{project_id}/backup/export")
async def backup_export_project(project_id: str, user: dict = Depends(get_current_user)):
    """Download a ZIP backup of a single project (data + its photos)."""
    await get_owned_project(project_id, user)
    data = await backup_mod.gather_user_data(user, project_ids=[project_id])
    photo_paths = backup_mod._collect_photo_paths(data)
    proj_label = data["projects"][0].get("name", "proyek") if data["projects"] else "proyek"
    zip_path, stats = await run_in_threadpool(
        backup_mod._assemble_zip, data, user, photo_paths, proj_label
    )
    slug = slugify(proj_label) or "proyek"
    filename = f"profinance-{slug}-{stats['stamp']}.zip"

    def _cleanup(p=zip_path):
        try:
            os.unlink(p)
        except Exception:
            pass

    return FileResponse(
        zip_path, media_type="application/zip", filename=filename,
        background=BackgroundTask(_cleanup),
    )


@api.post("/backup/run")
async def backup_run(user: dict = Depends(get_current_user)):
    """Create a stored backup now (saved to Object Storage), same as the weekly job."""
    doc = await backup_mod.create_stored_backup(user, kind="manual")
    if not doc:
        raise HTTPException(status_code=400, detail="Belum ada proyek untuk dibackup.")
    return doc


@api.get("/backups")
async def list_backups(user: dict = Depends(get_current_user)):
    """List stored backups (auto + manual) for the current user, newest first."""
    docs = await db.backups.find(
        {"user_id": user["user_id"]}, {"_id": 0, "storage_path": 0}
    ).sort("createdAt", -1).to_list(100)
    return docs


@api.get("/backups/{backup_id}/download")
async def download_backup(backup_id: str, request: Request, auth: Optional[str] = Query(None)):
    """Download a previously stored backup ZIP."""
    user = await _user_from_request_or_query(request, auth)
    doc = await db.backups.find_one({"id": backup_id, "user_id": user["user_id"]}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Backup tidak ditemukan")
    data, _ct = await run_in_threadpool(get_object, doc["storage_path"])
    return Response(
        content=data, media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{doc["filename"]}"'},
    )


@api.post("/backup/restore")
async def backup_restore(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    """Restore from an uploaded backup ZIP. Restores projects that don't already exist."""
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="File kosong.")
    if len(content) > 200 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File terlalu besar (maks 200MB).")
    try:
        result = await backup_mod.restore_from_zip(user, content)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        logger.error(f"restore failed: {e}")
        raise HTTPException(status_code=500, detail="Gagal memulihkan data.")
    return result


@api.post("/cron/weekly-backup")
async def cron_weekly_backup(request: Request, background: BackgroundTasks):
    # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
    secret = os.environ.get("WEBHOOK_CRON_SECRET")
    auth = request.headers.get("Authorization", "")
    token = auth[7:] if auth.startswith("Bearer ") else ""
    if not secret or not hmac.compare_digest(token, secret):
        raise HTTPException(status_code=401, detail="Unauthorized")
    background.add_task(backup_mod.run_weekly_backups)
    return {"ok": True, "queued": True}


@api.get("/")
async def root():
    return {"message": "ProFinance Interior API"}


@api.get("/categories")
async def categories():
    return {"income": INCOME_CATEGORIES, "expense": EXPENSE_CATEGORIES}


@api.post("/seed-demo")
async def seed_demo(user: dict = Depends(get_current_user)):
    existing = await db.projects.count_documents({"user_id": user["user_id"]})
    if existing > 0:
        raise HTTPException(status_code=400, detail="Data sudah ada, seed dilewati")

    def d(days_ago):
        return (datetime.now(timezone.utc) - timedelta(days=days_ago)).isoformat()

    thumbs = [
        "https://images.unsplash.com/photo-1633110187937-6e3b2f36dfca?crop=entropy&cs=srgb&fm=jpg&ixid=M3w4NTYxODd8MHwxfHNlYXJjaHwxfHxtb2Rlcm4lMjBsdXh1cnklMjBpbnRlcmlvciUyMGFyY2hpdGVjdHVyZSUyMGxpdmluZyUyMHJvb20lMjBraXRjaGVuJTIwb2ZmaWNlfGVufDB8fHx8MTc8OTgwMjMxN3ww&ixlib=rb-4.1.0&q=85",
        "https://images.pexels.com/photos/8089172/pexels-photo-8089172.jpeg?auto=compress&cs=tinysrgb&dpr=2&h=650&w=940",
        "https://images.unsplash.com/photo-1768321917661-d4f1a89d2185?crop=entropy&cs=srgb&fm=jpg&ixid=M3w3NDQ2MzR8MHwxfHNlYXJjaHwzfHxpbnRlcmlvciUyMGRlc2lnbiUyMGFyY2hpdGVjdHVyZSUyMGNvbnN0cnVjdGlvbiUyMHNpdGUlMjBmaW5pc2glMjByb29tfGVufDB8fHx8MTc4OTgwMjMxMHww&ixlib=rb-4.1.0&q=85",
    ]

    specs = [
        {
            "name": "Kitchen Set & Wall Panel Villa Canggu", "owner": "Bpk. Andika", "nominal": 285000000,
            "companyName": "Andika Family", "alamatProyek": "Canggu, Bali", "category": "Residensial",
            "start": 60, "end": -30, "thumb": 0,
            "tx": [
                ("in", 85000000, "Downpayment", "DP 30% kontrak", 58),
                ("in", 100000000, "Termin", "Termin 2 progress 50%", 30),
                ("out", 62000000, "Material", "HPL, plywood, hardware", 55),
                ("out", 18000000, "Material", "Granit & kaca backsplash", 40),
                ("out", 3200000, "Toll", "Ongkir & tol material", 38),
                ("out", 1500000, "Makan", "Konsumsi tukang mingguan", 20),
            ],
            "workers": [("Pak Slamet (Ketua)", 45000000), ("Tim Finishing HPL", 28000000)],
            "kasbon": [(0, 25000000), (1, 15000000)],
            "items": [("Pekerjaan Kabinet Bawah", 90000000, [(45, 100), (30, 100)]),
                      ("Wall Panel & Backdrop", 70000000, [(40, 60), (20, 85)]),
                      ("Finishing HPL & Hardware", 55000000, [(20, 40)])],
        },
        {
            "name": "Fitout Office Tower SCBD Lt.21", "owner": "PT Meridian Capital", "nominal": 620000000,
            "companyName": "PT Meridian Capital", "alamatProyek": "SCBD, Jakarta Selatan", "category": "Kantor",
            "start": 45, "end": -20, "thumb": 1,
            "tx": [
                ("in", 186000000, "Downpayment", "DP 30%", 44),
                ("in", 150000000, "Termin", "Termin partisi & plafon", 20),
                ("out", 210000000, "Material", "Partisi gypsum, kaca, plafon", 42),
                ("out", 45000000, "Material", "Karpet & vinyl flooring", 25),
                ("out", 4500000, "Bensin", "Operasional kendaraan", 18),
            ],
            "workers": [("Tim Partisi CV Jaya", 85000000), ("Tim Elektrikal", 42000000)],
            "kasbon": [(0, 40000000), (1, 20000000)],
            "items": [("Partisi & Dinding Gypsum", 220000000, [(40, 80), (20, 100)]),
                      ("Plafon & Pencahayaan", 150000000, [(30, 55)]),
                      ("Flooring Vinyl", 120000000, [(15, 30)])],
        },
        {
            "name": "Renovasi Rumah Cluster Bintaro", "owner": "Ibu Sari", "nominal": 140000000,
            "companyName": "-", "alamatProyek": "Bintaro Sektor 9", "category": "Residensial",
            "start": 30, "end": -15, "thumb": 2,
            "tx": [
                ("in", 42000000, "Downpayment", "DP awal", 28),
                ("out", 68000000, "Material", "Keramik, cat, plafon", 25),
                ("out", 12000000, "Material", "Pintu & kusen aluminium", 12),
                ("out", 900000, "Makan", "Konsumsi", 8),
            ],
            "workers": [("Pak Joko", 32000000)],
            "kasbon": [(0, 20000000)],
            "items": [("Pekerjaan Dinding & Cat", 55000000, [(25, 70)]),
                      ("Plafon & Listrik", 45000000, [(20, 40)])],
        },
    ]

    created = 0
    photo_i = 0
    for spec in specs:
        pid = str(uuid.uuid4())
        await db.projects.insert_one({
            "id": pid, "user_id": user["user_id"], "name": spec["name"], "owner": spec["owner"],
            "nominal": spec["nominal"], "companyName": spec["companyName"], "alamatProyek": spec["alamatProyek"],
            "tanggalMulai": d(spec["start"]), "targetSelesai": d(spec["end"]), "category": spec["category"],
            "status": "Berjalan", "thumbnail": thumbs[spec["thumb"]], "createdAt": d(spec["start"]),
            # Stable portal slug so client links survive the daily demo reset
            "portalSlug": f"{slugify(spec['name'])}-demo",
        })
        for typ, amt, cat, desc, days in spec["tx"]:
            await db.transactions.insert_one({
                "id": str(uuid.uuid4()), "project_id": pid, "type": typ, "amount": amt,
                "description": desc, "date": d(days), "category": cat, "receiptUrl": None, "createdAt": d(days),
            })
        workers = []
        for wname, boron in spec["workers"]:
            wid = str(uuid.uuid4())
            await db.workers.insert_one({"id": wid, "project_id": pid, "name": wname, "borongan": boron, "createdAt": d(spec["start"])})
            workers.append((wid, wname))
        for wi_idx, amt in spec["kasbon"]:
            wid, wname = workers[wi_idx]
            await db.transactions.insert_one({
                "id": str(uuid.uuid4()), "project_id": pid, "type": "out", "amount": amt,
                "description": f"Kasbon {wname}", "date": d(spec["start"] - 5),
                "category": f"Kasbon Tukang {wname}", "receiptUrl": None, "createdAt": d(spec["start"] - 5),
            })
        for iname, nilai, entries in spec["items"]:
            iid = str(uuid.uuid4())
            await db.work_items.insert_one({"id": iid, "project_id": pid, "name": iname, "nilai": nilai, "createdAt": d(spec["start"])})
            for days, prog in entries:
                pics = [DEMO_PHOTOS[photo_i % len(DEMO_PHOTOS)], DEMO_PHOTOS[(photo_i + 1) % len(DEMO_PHOTOS)]]
                photo_i += 2
                await db.progress_entries.insert_one({
                    "id": str(uuid.uuid4()), "workItemId": iid, "date": d(days), "progress": prog,
                    "notes": f"Update lapangan progres {prog}%", "photoUrls": pics, "createdAt": d(days),
                })
        created += 1

    return {"created": created}


app.include_router(api)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get("CORS_ORIGINS", "*").split(","),
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def startup():
    try:
        init_storage()
        logger.info("Storage initialized")
    except Exception as e:
        logger.error(f"Storage init failed: {e}")
    try:
        await db.orders.create_index("order_id", unique=True)
        await db.orders.create_index("user_id")
    except Exception as e:
        logger.error(f"Order index init failed: {e}")
    try:
        await db.backups.create_index("user_id")
        await db.backups.create_index([("user_id", 1), ("createdAt", -1)])
    except Exception as e:
        logger.error(f"Backup index init failed: {e}")
    try:
        await _seed_work_categories()
    except Exception as e:
        logger.error(f"Work categories seed failed: {e}")


@app.on_event("shutdown")
async def shutdown():
    auth_mod._client.close()

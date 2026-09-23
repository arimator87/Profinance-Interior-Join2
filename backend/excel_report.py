"""
Excel report builder for ProFinance Interior.

build_project_recap_xlsx(project, summary, recap, invoices, rab_computed, company)
-> bytes of a nicely formatted .xlsx with 3 sheets:
   1. Rekap Penagihan  (billing recap + project header)
   2. Daftar Invoice   (all invoices with computed columns)
   3. RAB              (bill of quantities, only if the project has one)
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---- Palette (matches the app's amber/slate brand) ----
BRAND = "B45309"        # amber-700
BRAND_LIGHT = "FEF3C7"  # amber-100
SLATE_DARK = "0F172A"   # slate-900
SLATE_HEAD = "1E293B"   # slate-800
SLATE_SOFT = "F1F5F9"   # slate-100
WHITE = "FFFFFF"
GREEN = "16A34A"
BLUE = "2563EB"
ORANGE = "EA580C"
VIOLET = "7C3AED"

RP_FMT = '"Rp"#,##0'
THIN = Side(style="thin", color="E2E8F0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _fmt_date(v):
    if not v:
        return "-"
    try:
        s = str(v)
        return datetime.fromisoformat(s.replace("Z", "+00:00")).strftime("%d %b %Y")
    except Exception:
        return str(v)[:10]


def _title(ws, text, subtitle, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c.fill = PatternFill("solid", fgColor=BRAND)
    ws.row_dimensions[1].height = 30
    for col in range(2, ncols + 1):
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=BRAND)
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        s = ws.cell(row=2, column=1, value=subtitle)
        s.font = Font(name="Calibri", size=10, italic=True, color=SLATE_HEAD)
        s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[2].height = 18


def _header_row(ws, row, headers, widths=None, fill=SLATE_HEAD):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, color=WHITE, size=10)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 22


def _money(cell):
    cell.number_format = RP_FMT
    cell.alignment = Alignment(horizontal="right")


def _build_recap_sheet(ws, project, summary, recap, company):
    _title(ws, "REKAP PENAGIHAN PROYEK",
            (company.get("companyName") or "") if company else "", 2)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 32

    row = 4
    info = [
        ("Nama Proyek", project.get("name", "-")),
        ("Klien", project.get("owner") or "-"),
        ("Status", project.get("status", "-")),
        ("Tanggal Cetak", datetime.now().strftime("%d %b %Y")),
    ]
    for label, val in info:
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = Font(bold=True, color=SLATE_HEAD, size=10)
        vc = ws.cell(row=row, column=2, value=val)
        vc.alignment = Alignment(horizontal="left")
        row += 1

    row += 1
    hc = ws.cell(row=row, column=1, value="RINGKASAN KEUANGAN")
    hc.font = Font(bold=True, color=WHITE)
    hc.fill = PatternFill("solid", fgColor=SLATE_HEAD)
    ws.cell(row=row, column=2).fill = PatternFill("solid", fgColor=SLATE_HEAD)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    hc.alignment = Alignment(horizontal="center")
    row += 1

    rows = [
        ("Nilai Kontrak", recap.get("nominal", 0), SLATE_DARK),
        ("Total Tertagih (invoice terkirim/lunas)", recap.get("totalBilled", 0), BLUE),
        ("Terbayar", recap.get("paid", 0), GREEN),
        ("Piutang (belum dibayar)", recap.get("receivable", 0), ORANGE),
        ("Nilai Draft (belum ditagihkan)", recap.get("draftAmount", 0), SLATE_HEAD),
        ("Retensi Ditahan", recap.get("retentionHeld", 0), VIOLET),
        ("Sisa Tagihan terhadap Kontrak", recap.get("sisaTagihan", 0), SLATE_DARK),
    ]
    for label, val, color in rows:
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = Font(color=SLATE_HEAD, size=10)
        lc.border = BORDER
        lc.fill = PatternFill("solid", fgColor=SLATE_SOFT)
        vc = ws.cell(row=row, column=2, value=val)
        vc.font = Font(bold=True, color=color, size=11)
        vc.border = BORDER
        _money(vc)
        row += 1

    row += 1
    note = ws.cell(row=row, column=1,
                   value=f"Jumlah invoice: {recap.get('invoiceCount', 0)}  |  "
                         f"Terkirim/Lunas: {recap.get('billedCount', 0)}")
    note.font = Font(italic=True, size=9, color=SLATE_HEAD)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)


TYPE_LABEL = {"proforma": "Proforma", "final": "Invoice", "retention": "Retensi"}


def _build_invoice_sheet(ws, invoices):
    _title(ws, "DAFTAR INVOICE", "", 12)
    headers = ["No", "Tipe", "No. Invoice", "No. Quotation", "Tanggal", "Jatuh Tempo",
               "Status", "Klien", "Subtotal", "PPN", "Retensi", "Total Tagihan"]
    widths = [5, 11, 20, 18, 13, 13, 11, 24, 15, 13, 13, 16]
    hr = 4
    _header_row(ws, hr, headers, widths)

    r = hr + 1
    tot_sub = tot_ppn = tot_ret = tot_due = 0
    for idx, inv in enumerate(invoices, start=1):
        c = inv.get("computed", {}) or {}
        tot_sub += c.get("subtotal", 0)
        tot_ppn += c.get("ppnAmount", 0)
        tot_ret += c.get("retentionAmount", 0)
        tot_due += c.get("amountDue", 0)
        vals = [
            idx,
            TYPE_LABEL.get(inv.get("type"), inv.get("type", "")),
            inv.get("number", ""),
            inv.get("quotationNo", "") or "-",
            _fmt_date(inv.get("invoiceDate")),
            _fmt_date(inv.get("dueDate")),
            inv.get("status", "Draft"),
            inv.get("clientName", "") or "-",
            c.get("subtotal", 0),
            c.get("ppnAmount", 0),
            c.get("retentionAmount", 0),
            c.get("amountDue", 0),
        ]
        for col, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.border = BORDER
            cell.font = Font(size=10)
            if col in (1, 2, 5, 6, 7):
                cell.alignment = Alignment(horizontal="center")
            if col >= 9:
                _money(cell)
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="FBFCFE")
        r += 1

    # Totals row
    tcell = ws.cell(row=r, column=8, value="TOTAL")
    tcell.font = Font(bold=True, color=WHITE)
    tcell.alignment = Alignment(horizontal="right")
    tcell.fill = PatternFill("solid", fgColor=BRAND)
    for col, val in ((9, tot_sub), (10, tot_ppn), (11, tot_ret), (12, tot_due)):
        cell = ws.cell(row=r, column=col, value=val)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BRAND)
        _money(cell)
    ws.freeze_panes = "A5"


def _build_rab_sheet(ws, rab):
    _title(ws, "RENCANA ANGGARAN BIAYA (RAB)", "", 6)
    headers = ["No", "Uraian Pekerjaan", "Volume", "Satuan", "Harga Satuan", "Jumlah"]
    widths = [5, 44, 10, 10, 18, 20]
    hr = 4
    _header_row(ws, hr, headers, widths)

    r = hr + 1
    for si, sec in enumerate(rab.get("sections", []) or [], start=1):
        # Section header row
        sc = ws.cell(row=r, column=1, value=_roman(si))
        sc.alignment = Alignment(horizontal="center")
        nc = ws.cell(row=r, column=2, value=(sec.get("name") or "").upper())
        for col in range(1, 7):
            cell = ws.cell(row=r, column=col)
            cell.fill = PatternFill("solid", fgColor=BRAND_LIGHT)
            cell.font = Font(bold=True, color=SLATE_DARK, size=10)
            cell.border = BORDER
        subc = ws.cell(row=r, column=6, value=sec.get("subtotal", 0))
        subc.font = Font(bold=True, color=SLATE_DARK, size=10)
        _money(subc)
        r += 1

        for j, sub in enumerate(sec.get("subItems", []) or [], start=1):
            vals = [
                f"{si}.{j}",
                sub.get("name", ""),
                sub.get("qty", 0),
                sub.get("unit", "Unit"),
                sub.get("hargaSatuan", 0),
                sub.get("nilai", 0),
            ]
            for col, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=col, value=v)
                cell.border = BORDER
                cell.font = Font(size=10)
                if col == 1:
                    cell.alignment = Alignment(horizontal="center")
                elif col in (3, 4):
                    cell.alignment = Alignment(horizontal="center")
                elif col in (5, 6):
                    _money(cell)
            r += 1
            # material breakdown as light sub-lines
            for m in sub.get("materials", []) or []:
                mc = ws.cell(row=r, column=2, value=f"   • {m.get('name', '')}")
                mc.font = Font(size=9, italic=True, color=SLATE_HEAD)
                mv = ws.cell(row=r, column=6, value=m.get("nilai", 0))
                mv.font = Font(size=9, italic=True, color=SLATE_HEAD)
                _money(mv)
                for col in range(1, 7):
                    ws.cell(row=r, column=col).border = BORDER
                r += 1

    r += 1
    summary_rows = [("Total Item Pekerjaan", rab.get("totalItems", 0), False)]
    if rab.get("discount", 0):
        summary_rows.append(("Diskon", -rab.get("discount", 0), False))
        summary_rows.append(("Setelah Diskon", rab.get("afterDiscount", 0), False))
    if rab.get("ppnEnabled"):
        summary_rows.append((f"PPN {rab.get('ppnPercent', 0):g}%", rab.get("ppnAmount", 0), False))
    summary_rows.append(("GRAND TOTAL", rab.get("grandTotal", 0), True))

    for label, val, is_grand in summary_rows:
        lc = ws.cell(row=r, column=5, value=label)
        vc = ws.cell(row=r, column=6, value=val)
        if is_grand:
            lc.font = Font(bold=True, color=WHITE)
            vc.font = Font(bold=True, color=WHITE)
            lc.fill = PatternFill("solid", fgColor=BRAND)
            vc.fill = PatternFill("solid", fgColor=BRAND)
        else:
            lc.font = Font(bold=True, color=SLATE_HEAD, size=10)
            vc.font = Font(color=SLATE_HEAD, size=10)
        lc.alignment = Alignment(horizontal="right")
        _money(vc)
        r += 1

    termins = rab.get("termins", []) or []
    if termins:
        r += 1
        th = ws.cell(row=r, column=2, value="TERMIN PEMBAYARAN")
        th.font = Font(bold=True, color=SLATE_HEAD, size=10)
        r += 1
        for t in termins:
            ws.cell(row=r, column=2, value=t.get("label", "")).font = Font(size=10)
            pc = ws.cell(row=r, column=5, value=f"{t.get('percent', 0):g}%")
            pc.alignment = Alignment(horizontal="right")
            pc.font = Font(size=10)
            nc = ws.cell(row=r, column=6, value=t.get("nominal", 0))
            nc.font = Font(size=10)
            _money(nc)
            r += 1


def _roman(n):
    vals = [(10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out


def build_project_recap_xlsx(project, summary, recap, invoices, rab_computed=None, company=None):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Rekap Penagihan"
    _build_recap_sheet(ws1, project, summary, recap, company or {})

    ws2 = wb.create_sheet("Daftar Invoice")
    _build_invoice_sheet(ws2, invoices or [])

    if rab_computed and (rab_computed.get("sections")):
        ws3 = wb.create_sheet("RAB")
        _build_rab_sheet(ws3, rab_computed)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
